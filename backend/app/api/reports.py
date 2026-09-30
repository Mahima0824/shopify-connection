from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import JSONResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session
from app.api.auth import get_current_user
from app.database import get_db

router = APIRouter(prefix="/api/v1/reports", tags=["reports"])

EXPORTABLE = ("dashboard", "sales", "orders", "payments", "refunds",
              "profit", "courier", "gst", "reconciliation")


def _err(status: int, code: str, message: str) -> JSONResponse:
    return JSONResponse(status_code=status,
                        content={"success": False, "error": {"code": code, "message": message}})


def _range(preset: str | None, from_val: str | None, to_val: str | None,
           db: Session, business_id: str):
    from app.services import report_service as rs
    from app.services import accounting_service as acct
    try:
        return rs.resolve_range(preset, from_val, to_val,
                                fy_start_month=acct._fy_start_month(db, business_id))
    except ValueError as e:
        raise HTTPException(400, str(e))


def _filters(status: str | None, courier: str | None, payment_method: str | None,
             customer: str | None, product: str | None, state: str | None) -> dict:
    return {"status": status, "courier": courier, "payment_method": payment_method,
            "customer": customer, "product": product, "state": state}


def _maybe_export(name: str, data: dict, fmt: str | None):
    if not fmt:
        return {"success": True, "data": data}
    from app.services import report_service as rs
    rows_fn = {
        "sales": lambda: ([["Order", "Date", "Financial", "Operational", "Total"]] +
                          [[r["name"], r["date"], r["financial"], r["operational"], r["total"]]
                           for r in data.get("rows", [])]),
        "orders": lambda: ([["Order", "Date", "Financial", "Operational", "Total"]] +
                           [[r["name"], r["date"], r["financial"], r["operational"], r["total"]]
                            for r in data.get("rows", [])]),
        "payments": lambda: ([["ID", "Order", "Amount", "Status", "Method"]] +
                             [[r["id"], r["order_id"], r["amount"], r["status"], r["method"]]
                              for r in data.get("rows", [])]),
        "refunds": lambda: ([["ID", "Order", "Amount", "Status"]] +
                            [[r["id"], r["order_id"], r["amount"], r["status"]]
                             for r in data.get("rows", [])]),
        "profit": lambda: ([["Metric", "Value"]] +
                           [["Net sales", data["revenue"]["net_exclusive"]],
                            ["GST", data["revenue"]["gst"]],
                            ["COGS", data["profit"]["cogs"]],
                            ["Gross profit", data["profit"]["gross_profit"]],
                            ["Operating profit", data["profit"]["operating_profit"]],
                            ["Margin %", data["profit"]["margin_pct"]],
                            ["Label", data["profit"]["label"]],
                            ["Warning", data["profit"]["warning"]]]),
        "courier": lambda: ([["AWB", "Carrier", "Status", "Order"]] +
                            [[r["awb"], r["carrier"], r["status"], r["order_id"]]
                             for r in data.get("rows", [])]),
        "gst": lambda: ([["Order", "Valid", "Jurisdiction", "Taxable", "CGST", "SGST", "IGST", "Total"]] +
                        [[r["order_name"], r["valid"], r["jurisdiction"],
                          r["invoice_check"]["taxable"], r["invoice_check"]["cgst"],
                          r["invoice_check"]["sgst"], r["invoice_check"]["igst"],
                          r["invoice_check"]["total"]] for r in data.get("rows", [])]),
        "reconciliation": lambda: ([["Order", "Code", "Severity", "Message", "Resolved"]] +
                                   [[r["order_id"], r["code"], r["severity"], r["message"], r["resolved"]]
                                    for r in data.get("rows", [])]),
        "dashboard": lambda: ([["KPI", "Value"]] +
                              [[k, v] for k, v in data.get("kpis", {}).items()]),
    }
    table = rows_fn[name]()
    headers, body = table[0], table[1:]
    # 1-based money/date columns per report so XLSX holds numerics + real dates (#79).
    _specs = {
        "sales": ((5,), (2,)), "orders": ((5,), (2,)),
        "payments": ((3,), ()), "refunds": ((3,), ()),
        "profit": ((2,), ()), "courier": ((), ()),
        "gst": ((4, 5, 6, 7, 8), ()),
        "reconciliation": ((), ()), "dashboard": ((2,), ()),
    }
    money_cols, date_cols = _specs[name]
    if len(body) > rs.ASYNC_EXPORT_ROW_LIMIT:
        return JSONResponse(status_code=202, content={
            "success": False,
            "error": {"code": "EXPORT_TOO_LARGE",
                      "message": f"{len(body)} rows exceed sync limit "
                                 f"({rs.ASYNC_EXPORT_ROW_LIMIT}); run async per #54/#78."}})
    if fmt == "csv":
        content = rs.export_table_csv(headers, body)
        return Response(content=content, media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{name}_report.csv"'})
    if fmt == "xlsx":
        content = rs.export_table_xlsx(name.title(), headers, body,
                                       money_cols=money_cols, date_cols=date_cols)
        return Response(
            content=content,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{name}_report.xlsx"'})
    return _err(400, "BAD_REQUEST", "format must be csv or xlsx")


def _endpoint(name: str, builder):
    def ep(preset: str | None = None,
           from_val: str | None = Query(default=None, alias="from"),
           to: str | None = None, status: str | None = None,
           courier: str | None = None, payment_method: str | None = None,
           customer: str | None = None, product: str | None = None,
           state: str | None = None, format: str | None = None,
           db: Session = Depends(get_db), u: dict = Depends(get_current_user)):
        s, e = _range(preset, from_val, to, db, u.get("business_id"))
        data = builder(db, u.get("business_id"), s, e,
                       **_filters(status, courier, payment_method, customer, product, state))
        return _maybe_export(name, data, format)
    return ep


def _wire():
    from app.services import report_service as rs
    for _name, _fn in (
            ("dashboard", rs.dashboard_report), ("sales", rs.sales_report),
            ("orders", rs.orders_report), ("payments", rs.payments_report),
            ("refunds", rs.refunds_report), ("profit", rs.profit_report),
            ("courier", rs.courier_report), ("gst", rs.gst_report),
            ("reconciliation", rs.reconciliation_report)):
        router.get(f"/{_name}")(_endpoint(_name, _fn))


_wire()


@router.get("/gst/validate")
def gst_validate(order_id: str, db: Session = Depends(get_db),
                 u: dict = Depends(get_current_user)):
    from app.services import report_service as rs
    try:
        return {"success": True, "data": rs.validate_order_gst(db, u.get("business_id"), order_id)}
    except ValueError as e:
        return _err(404 if "not found" in str(e).lower() else 400, "BAD_REQUEST", str(e))


@router.get("/presets")
def presets():
    from app.services.report_service import PRESETS
    return {"success": True, "data": {"presets": list(PRESETS)}}



class CostItem(BaseModel):
    key: str
    amount: float
    source: str = "MANUAL"
    effective_from: str
    note: str | None = None


class CostsIn(BaseModel):
    items: list[CostItem]


@router.get("/monthly")
def monthly(month: str, db: Session = Depends(get_db), u: dict = Depends(get_current_user)):
    from app.services.report_service import monthly_report, month_window
    try:
        month_window(month)
    except (ValueError, AttributeError):
        raise HTTPException(400, "month must be YYYY-MM")
    return {"success": True, "data": monthly_report(db, u.get("business_id"), month)}


@router.get("/costs")
def get_costs(db: Session = Depends(get_db), u: dict = Depends(get_current_user)):
    from datetime import datetime, timezone
    from app.models.cost import CostRule
    from app.services.cost_service import KEYS
    now = datetime.now(timezone.utc)
    out = []
    for key in KEYS:
        r = db.query(CostRule).filter_by(business_id=u.get("business_id"), key=key).filter(
            CostRule.effective_to.is_(None)).first()
        out.append({"key": key, "amount": float(r.amount or 0) if r else 0.0,
                    "source": r.source if r else "DEFAULT"})
    return {"success": True, "data": {"items": out}}


@router.put("/costs")
def put_costs(body: CostsIn, db: Session = Depends(get_db), u: dict = Depends(get_current_user)):
    from datetime import datetime
    from app.services.cost_service import set_cost, KEYS
    if u.get("role") not in ("ADMIN", "ACCOUNTANT"):
        raise HTTPException(403, "Accountant role required")
    for item in body.items:
        if item.key not in KEYS:
            raise HTTPException(400, f"Unknown cost key '{item.key}'.")
        try:
            ef = datetime.fromisoformat(item.effective_from)
        except ValueError:
            raise HTTPException(400, f"Invalid effective_from for {item.key}.")
        set_cost(db, u.get("business_id"), item.key, item.amount, item.source, ef, item.note)
    db.commit()
    return {"success": True, "data": {"updated": len(body.items)}}


@router.get("/monthly/export")
def monthly_export(month: str, db: Session = Depends(get_db), u: dict = Depends(get_current_user)):
    from app.services.report_service import monthly_report, month_window
    try:
        month_window(month)
    except (ValueError, AttributeError):
        raise HTTPException(400, "month must be YYYY-MM")
    data = monthly_report(db, u.get("business_id"), month)
    content = _workbook(data)
    headers = {"Content-Disposition": f'attachment; filename="monthly_report_{month}.xlsx"'}
    return Response(content=content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers=headers)


def _workbook(data: dict) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E5F")
    money_fmt = "#,##0.00"

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name) if name != "Summary" else wb.active
        if name == "Summary":
            ws.title = "Summary"
        ws.freeze_panes = "A2"
        ws.append(headers)
        for c in ws[1]:
            c.font = hdr_font
            c.fill = hdr_fill
        ws.auto_filter.ref = ws.dimensions
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            w = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(w + 2, 40)
        return ws

    p = data["profitability"]
    sheet("Summary", ["Metric", "Value"], [
        ["Month", data["month"]], ["Generated", data["generated_at"]],
        ["Orders", data["orders"]["total"]], ["Gross", data["money"]["gross"]],
        ["Net revenue", p["net_revenue"]], ["Costs", p["cost_total"]],
        [p["label"], p["profit"]]])
    sheet("Orders", ["Name", "Financial", "Operational", "Total"],
          [[r["name"], r["financial"], r["operational"], r["total"]] for r in data.get("order_rows", [])])
    sheet("Shipments", ["AWB", "Carrier", "Status", "Location"],
          [[r["awb"], r["carrier"], r["status"], r["location"]] for r in data.get("shipment_rows", [])])
    m = data["money"]
    sheet("Money", ["Metric", "Value"],
          [["Gross", m["gross"]], ["Discounts", m["discounts"]], ["Refunds", m["refunds"]],
           ["Expected", m["expected"]], ["Collected", m["collected"]], ["Settled", m["settled"]],
           ["Pending", m["pending"]], ["Fees", m["fees"]], ["Net", m["net"]]])
    sheet("Exceptions", ["Metric", "Value"],
          [["Open", data["exceptions"]["open"]]] +
          [[f"Severity {k}", v] for k, v in data["exceptions"]["by_severity"].items()])
    sheet("Profitability", ["Line", "Total", "Source"],
          [["Gross", p["gross"], ""], ["Discounts", p["discounts"], ""], ["Refunds", p["refunds"], ""],
           ["Net revenue", p["net_revenue"], ""]] +
          [[c["key"], c["total"], c["source"]] for c in data["costs"]["lines"]] +
          [[p["label"], p["profit"], ""]])
    return _finalize(wb, money_fmt)


def _finalize(wb, money_fmt: str) -> bytes:
    import io
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = money_fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
