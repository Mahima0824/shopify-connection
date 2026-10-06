import csv
import hashlib
import io
from datetime import datetime, timezone

TYPES = ("COURIER_SETTLEMENT", "BANK_STATEMENT", "PAYMENT_GATEWAY_STATEMENT",
         "COURIER_SHIPMENT_REPORT", "SHOPIFY_ORDER_EXPORT")

ALIASES = {
    "awb_number": ["awb no", "awb", "awb number", "tracking id", "tracking number"],
    "external_reference": ["txn ref", "transaction ref", "transaction id", "settlement id",
                           "settlement reference"],
    "reference_number": ["utr", "utr number", "utr no", "bank reference", "bank ref", "ref no",
                         "reference", "reference number", "transaction ref"],
    "description": ["description", "narration", "particulars", "remarks", "details"],
    "order_reference": ["order id", "order", "order name", "order number", "shopify order"],
    "transaction_date": ["date", "txn date", "transaction date", "posting date", "settlement date"],
    "value_date": ["value date", "effective date"],
    "gross_amount": ["cod amount", "gross", "amount", "credit"],
    "fee_amount": ["fee", "courier fee", "charges", "commission"],
    "net_amount": ["net remittance", "net", "settled amount", "settlement amount"],
    "debit": ["debit", "withdrawal", "withdrawals", "dr", "debit amount"],
    "credit": ["credit", "deposit", "deposits", "cr", "credit amount"],
    "balance": ["balance", "closing balance", "running balance", "available balance"],
    "transaction_type": ["type", "txn type"],
    "status": ["status", "txn status"],
}

REQUIRED = {
    "COURIER_SETTLEMENT": ["awb_number"],
    "BANK_STATEMENT": [],
    "PAYMENT_GATEWAY_STATEMENT": ["external_reference"],
    "COURIER_SHIPMENT_REPORT": ["awb_number"],
    "SHOPIFY_ORDER_EXPORT": ["order_reference"],
}

# Canonical bank columns for the Date/Description/Reference/Debit/Credit/Balance mapping (#34).
BANK_COLUMNS = ("transaction_date", "description", "reference_number", "debit", "credit", "balance")


def file_sha(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _norm_header(h: str) -> str:
    return " ".join((h or "").strip().lower().split())


def map_columns(headers: list[str], statement_type: str) -> dict:
    normed = {_norm_header(h): h for h in headers}
    mapping: dict[str, str] = {}
    for field, aliases in ALIASES.items():
        for a in aliases:
            if a in normed:
                mapping[field] = normed[a]
                break
    missing = [f for f in REQUIRED.get(statement_type, []) if f not in mapping]
    if missing:
        raise ValueError(f"MISSING_REQUIRED_COLUMN: {', '.join(missing)}")
    return mapping


def _rows_csv(content: bytes) -> tuple[list[str], list[dict]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return reader.fieldnames or [], list(reader)


def _rows_xlsx(content: bytes) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook
    wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h or "").strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({h: ("" if v is None else str(v)) for h, v in zip(headers, r)})
    return headers, out


def parse_statement(content: bytes, filename: str) -> tuple[list[str], list[dict]]:
    if filename.lower().endswith(".xlsx"):
        return _rows_xlsx(content)
    if filename.lower().endswith(".csv"):
        return _rows_csv(content)
    raise ValueError("INVALID_STATEMENT: only .csv and .xlsx are accepted.")


def _fnum(v) -> float:
    try:
        return float(str(v or 0).replace(",", "").strip() or 0)
    except ValueError:
        return 0.0


def _fdate(v):
    v = (str(v) if v is not None else "").strip()
    if not v:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S %z", "%Y-%m-%d %H:%M:%S"):
        try:
            dt = datetime.strptime(v, fmt)
            return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def row_to_fields(row: dict, mapping: dict) -> dict:
    g = lambda f: (row.get(mapping[f]) or "") if f in mapping else ""
    debit_v, credit_v = _fnum(g("debit")), _fnum(g("credit"))
    gross_v, net_v = _fnum(g("gross_amount")), _fnum(g("net_amount"))
    if "credit" in mapping and "gross_amount" not in mapping and gross_v == 0:
        gross_v = credit_v
    if "credit" in mapping and "net_amount" not in mapping and net_v == 0:
        net_v = credit_v - debit_v
    return {
        "external_reference": (g("external_reference") or "").strip() or None,
        "awb_number": (g("awb_number") or "").strip() or None,
        "order_reference": (g("order_reference") or "").strip() or None,
        "transaction_date": _fdate(g("transaction_date")),
        "value_date": _fdate(g("value_date")),
        "gross_amount": gross_v,
        "fee_amount": _fnum(g("fee_amount")),
        "net_amount": net_v,
        "description": (g("description") or "").strip() or None,
        "reference_number": (g("reference_number") or "").strip() or None,
        "debit": debit_v,
        "credit": credit_v,
        "balance": _fnum(g("balance")),
        "transaction_type": (g("transaction_type") or "").strip() or None,
        "status": (g("status") or "").strip() or None,
    }


def import_identity(business_id: str, statement_type: str, bank_account_id: str | None,
                    fields: dict) -> str:
    """Stable bank import identity per #71: same external event -> same identity, never twice."""
    parts = [
        business_id or "", statement_type or "", bank_account_id or "",
        (fields.get("reference_number") or "").strip().upper(),
        (fields.get("external_reference") or "").strip().upper(),
        (fields.get("awb_number") or "").strip().upper(),
        (fields.get("order_reference") or "").strip().upper(),
        str(fields.get("transaction_date") or ""),
        f"{float(fields.get('debit') or 0):.2f}",
        f"{float(fields.get('credit') or 0):.2f}",
        f"{float(fields.get('net_amount') or 0):.2f}",
        (fields.get("description") or "").strip().lower(),
    ]
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()


def match_row(db, business_id: str, row) -> tuple[str, str | None, str | None]:
    """Returns (result, order_id|None, shipment_id|None). Never matches on amount alone."""
    from app.models.shipment import Shipment
    from app.models.order import Order
    from app.models.sla import ShipmentFinancial
    from app.services.shipment_service import is_awaiting_awb
    if row.external_reference:
        fin = db.query(ShipmentFinancial).filter_by(
            business_id=business_id, settlement_reference=row.external_reference).first()
        if fin is not None:
            return "MATCHED", fin.order_id, fin.shipment_id
    if row.awb_number:
        s = db.query(Shipment).filter_by(business_id=business_id, awb_number=row.awb_number).first()
        # Guarded even though a real bank row cannot contain our placeholder:
        # this match decides that money is SETTLED against a shipment. Matching
        # an order that was never dispatched would be a false financial match,
        # and one clause is cheaper than proving the collision impossible. The
        # amount fallback below needs a ShipmentFinancial row, which an
        # un-pushed awaiting shipment cannot have, so it needs no guard.
        if s is not None and not is_awaiting_awb(s.awb_number):
            return "MATCHED", s.order_id, s.id
    if row.order_reference:
        ref = row.order_reference.strip()
        o = db.query(Order).filter_by(business_id=business_id, shopify_order_id=ref).first()
        if o is None:
            o = db.query(Order).filter_by(business_id=business_id, shopify_order_name=ref).first()
        if o is None and ref.startswith("#"):
            o = db.query(Order).filter_by(business_id=business_id, shopify_order_name=ref).first()
        if o is not None:
            return "MATCHED", o.id, None
    if row.transaction_date and row.net_amount:
        lo = row.transaction_date
        cands = []
        for s in db.query(Shipment).filter_by(business_id=business_id).all():
            fin = db.query(ShipmentFinancial).filter_by(shipment_id=s.id).first()
            if fin is None:
                continue
            try:
                exp = float(fin.net_settlement or 0)
            except (TypeError, ValueError):
                continue
            if abs(exp - float(row.net_amount or 0)) > 1.0:
                continue
            if s.last_checkpoint_at and abs((s.last_checkpoint_at - lo).days) <= 3:
                cands.append(s)
        if len(cands) == 1:
            return "PARTIALLY_MATCHED", cands[0].order_id, cands[0].id
    return "UNMATCHED", None, None
