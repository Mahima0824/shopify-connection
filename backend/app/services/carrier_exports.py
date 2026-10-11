"""Multi-carrier Excel bulk export service for Delhivery, Blue Dart, DTDC, and India Post."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.services.india_post_export import build_workbook as build_indiapost_workbook, _get, to_num_or_none, to_int_or_none

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BORDER_THIN = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)

def get_cod_collectable_amount(order) -> float:
    """Calculate the exact COD collectable amount:
    Customer pays ₹300 advance online.
    Remaining balance = Total - ₹300.
    COD handling charge = +₹50.
    Cash on Delivery to collect = Remaining balance + ₹50.
    """
    is_cod = str(_get(order, "cod_mode") or "").strip().upper() == "COD"
    if not is_cod:
        return 0.0
    total = float(to_num_or_none(_get(order, "total_amount")) or 0.0)
    if total <= 0:
        return float(to_num_or_none(_get(order, "cod_value")) or 0.0)
    remaining = max(0.0, total - 300.0)
    return round(remaining + 50.0, 2)

# -------------------------------------------------------------
# 1. DELHIVERY B2C BULK UPLOAD WORKBOOK
# -------------------------------------------------------------
DELHIVERY_HEADERS = [
    "Order Number",
    "Customer Name",
    "Customer Phone",
    "Customer Address 1",
    "Customer Address 2",
    "Customer City",
    "Customer State",
    "Customer Pincode",
    "Product Description",
    "Quantity",
    "Payment Mode",
    "Total Order Value",
    "COD Amount",
    "Dead Weight (kg)",
    "Length (cm)",
    "Width (cm)",
    "Height (cm)",
    "Pickup Location",
    "Return Address",
    "Return Pincode",
]

def build_delhivery_workbook(orders) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Delhivery_B2C_Orders"

    # Header styling (Delhivery Red / Dark Slate)
    header_fill = PatternFill(start_color="D32F2F", end_color="D32F2F", fill_type="solid")
    ws.append(DELHIVERY_HEADERS)

    for col_idx in range(1, len(DELHIVERY_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, o in enumerate(orders, start=2):
        name = _get(o, "receiver_name") or "Customer"
        mobile = _get(o, "receiver_mobile") or ""
        add1 = _get(o, "receiver_add1") or ""
        add2 = _get(o, "receiver_add2") or ""
        city = _get(o, "receiver_city") or ""
        state = _get(o, "receiver_state") or ""
        pincode = _get(o, "receiver_pincode") or ""

        order_no = _get(o, "shopify_order_name") or _get(o, "internal_order_number") or str(_get(o, "id"))
        is_cod = str(_get(o, "cod_mode") or "").upper() == "COD"
        total_amt = float(to_num_or_none(_get(o, "total_amount")) or 0)
        cod_amt = get_cod_collectable_amount(o)

        weight_kg = round(float(to_num_or_none(_get(o, "weight_grams")) or 930) / 1000.0, 2)
        l_cm = to_int_or_none(_get(o, "length_cm")) or 30
        w_cm = to_int_or_none(_get(o, "breadth_cm")) or 20
        h_cm = to_int_or_none(_get(o, "height_cm")) or 5

        # Product summary from line items if present
        items = _get(o, "items") or []
        if items and isinstance(items, list):
            item_desc = ", ".join([str(_get(it, "title") or _get(it, "sku") or "Saree") for it in items[:2]])
            qty = sum([int(_get(it, "quantity") or 1) for it in items])
        else:
            item_desc = "Saree / Silk Apparel"
            qty = 1

        row_data = [
            order_no,
            name,
            str(mobile),
            add1,
            add2,
            city,
            state,
            str(pincode),
            item_desc,
            qty,
            "COD" if is_cod else "Prepaid",
            total_amt,
            cod_amt,
            weight_kg,
            l_cm,
            w_cm,
            h_cm,
            "Surat Warehouse",
            "FF-138 Rajhans Imperia, Ring Road, Surat, Gujarat",
            "395002",
        ]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    return wb


# -------------------------------------------------------------
# 2. BLUE DART BULK UPLOAD WORKBOOK
# -------------------------------------------------------------
BLUEDART_HEADERS = [
    "Order No / Reference",
    "Consignee Name",
    "Consignee Mobile",
    "Address Line 1",
    "Address Line 2",
    "City",
    "State",
    "Pincode",
    "Product Code",
    "Payment Mode",
    "Declared Value",
    "Collectable Amount (COD)",
    "Weight in Kg",
    "Pieces",
    "Commodity Description",
    "Sender Name",
    "Sender Phone",
    "Sender Address",
    "Sender Pincode",
]

def build_bluedart_workbook(orders) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "BlueDart_Manifest"

    # Blue Dart Deep Blue
    header_fill = PatternFill(start_color="0D47A1", end_color="0D47A1", fill_type="solid")
    ws.append(BLUEDART_HEADERS)

    for col_idx in range(1, len(BLUEDART_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, o in enumerate(orders, start=2):
        name = _get(o, "receiver_name") or "Consignee"
        mobile = _get(o, "receiver_mobile") or ""
        add1 = _get(o, "receiver_add1") or ""
        add2 = _get(o, "receiver_add2") or ""
        city = _get(o, "receiver_city") or ""
        state = _get(o, "receiver_state") or ""
        pincode = _get(o, "receiver_pincode") or ""

        order_no = _get(o, "shopify_order_name") or _get(o, "internal_order_number") or str(_get(o, "id"))
        is_cod = str(_get(o, "cod_mode") or "").upper() == "COD"
        total_amt = float(to_num_or_none(_get(o, "total_amount")) or 0)
        cod_amt = get_cod_collectable_amount(o)

        weight_kg = round(float(to_num_or_none(_get(o, "weight_grams")) or 930) / 1000.0, 2)
        items = _get(o, "items") or []
        qty = sum([int(_get(it, "quantity") or 1) for it in items]) if items and isinstance(items, list) else 1

        row_data = [
            order_no,
            name,
            str(mobile),
            add1,
            add2,
            city,
            state,
            str(pincode),
            "E-COMM",
            "COD" if is_cod else "NON-COD",
            total_amt,
            cod_amt,
            weight_kg,
            qty,
            "Saree - Designer Indian Wear",
            "Reshamgath",
            "9016822651",
            "Rajhans Imperia, Ring Road, Surat",
            "395002",
        ]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    return wb


# -------------------------------------------------------------
# 3. DTDC BULK CONSIGNMENT WORKBOOK
# -------------------------------------------------------------
DTDC_HEADERS = [
    "Reference Number",
    "Customer Name",
    "Mobile Number",
    "Address Line 1",
    "Address Line 2",
    "Pincode",
    "City",
    "State",
    "Service Type",
    "Payment Mode",
    "COD Amount",
    "Declared Value",
    "Weight (Kg)",
    "Pieces",
    "Commodity Description",
    "Length (cm)",
    "Width (cm)",
    "Height (cm)",
]

def build_dtdc_workbook(orders) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "DTDC_Consignments"

    # DTDC Crimson / Red
    header_fill = PatternFill(start_color="B71C1C", end_color="B71C1C", fill_type="solid")
    ws.append(DTDC_HEADERS)

    for col_idx in range(1, len(DTDC_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, o in enumerate(orders, start=2):
        name = _get(o, "receiver_name") or "Customer"
        mobile = _get(o, "receiver_mobile") or ""
        add1 = _get(o, "receiver_add1") or ""
        add2 = _get(o, "receiver_add2") or ""
        city = _get(o, "receiver_city") or ""
        state = _get(o, "receiver_state") or ""
        pincode = _get(o, "receiver_pincode") or ""

        order_no = _get(o, "shopify_order_name") or _get(o, "internal_order_number") or str(_get(o, "id"))
        is_cod = str(_get(o, "cod_mode") or "").upper() == "COD"
        total_amt = float(to_num_or_none(_get(o, "total_amount")) or 0)
        cod_amt = get_cod_collectable_amount(o)

        weight_kg = round(float(to_num_or_none(_get(o, "weight_grams")) or 930) / 1000.0, 2)
        l_cm = to_int_or_none(_get(o, "length_cm")) or 30
        w_cm = to_int_or_none(_get(o, "breadth_cm")) or 20
        h_cm = to_int_or_none(_get(o, "height_cm")) or 5

        items = _get(o, "items") or []
        qty = sum([int(_get(it, "quantity") or 1) for it in items]) if items and isinstance(items, list) else 1

        row_data = [
            order_no,
            name,
            str(mobile),
            add1,
            add2,
            str(pincode),
            city,
            state,
            "B2C EXPRESS",
            "COD" if is_cod else "PREPAID",
            cod_amt,
            total_amt,
            weight_kg,
            qty,
            "Saree / Apparel",
            l_cm,
            w_cm,
            h_cm,
        ]
        ws.append(row_data)

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = REGULAR_FONT
            cell.border = BORDER_THIN

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    return wb
