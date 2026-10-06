from datetime import datetime, timezone
from calendar import monthrange
from decimal import Decimal, ROUND_HALF_UP

DISPLAY_TZ = "Asia/Kolkata"
INDIA_FY_START_MONTH = 4  # #27/#77: 1 April -> 31 March

PRESETS = ("today", "yesterday", "this_week", "last_week", "this_month",
           "last_month", "this_quarter", "this_year", "financial_year",
           "last_fy", "custom")


def _utc(dt) -> datetime:
    if dt is None:
        return datetime.now(timezone.utc)
    if isinstance(dt, str):
        dt = datetime.fromisoformat(dt)
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def _d(v) -> Decimal:
    try:
        return Decimal(str(v or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except Exception:
        return Decimal("0.00")


def resolve_preset(preset: str, now: datetime | None = None,
                   fy_start_month: int = INDIA_FY_START_MONTH,
                   custom_from: str | None = None,
                   custom_to: str | None = None) -> tuple[datetime, datetime]:
    """Global date presets per #27. Returns UTC-aware [start, end)."""
    from datetime import timedelta
    d = _utc(now) if now else datetime.now(timezone.utc)
    day = datetime(d.year, d.month, d.day, tzinfo=timezone.utc)
    p = (preset or "").lower()
    if p == "today":
        return day, day + timedelta(days=1)
    if p == "yesterday":
        return day - timedelta(days=1), day
    if p == "this_week":  # Monday..Sunday
        start = day - timedelta(days=day.weekday())
        return start, start + timedelta(days=7)
    if p == "last_week":
        start = day - timedelta(days=day.weekday() + 7)
        return start, start + timedelta(days=7)
    if p == "this_month":
        return month_window(f"{d.year}-{d.month:02d}")
    if p == "last_month":
        y, m = (d.year, d.month - 1) if d.month > 1 else (d.year - 1, 12)
        return month_window(f"{y}-{m:02d}")
    if p == "this_quarter":
        q = (d.month - 1) // 3
        sm = q * 3 + 1
        em = sm + 3
        ey = d.year + (1 if em > 12 else 0)
        em = em if em <= 12 else em - 12
        return (datetime(d.year, sm, 1, tzinfo=timezone.utc),
                datetime(ey, em, 1, tzinfo=timezone.utc))
    if p == "this_year":
        return datetime(d.year, 1, 1, tzinfo=timezone.utc), datetime(d.year + 1, 1, 1, tzinfo=timezone.utc)
    if p == "financial_year":
        from app.services import ledger_service as _ls
        return _ls.fy_bounds(d, fy_start_month=fy_start_month)
    if p == "last_fy":
        from app.services import ledger_service as _ls
        s, _e = _ls.fy_bounds(d, fy_start_month=fy_start_month)
        prev = s - timedelta(days=1)
        return _ls.fy_bounds(prev, fy_start_month=fy_start_month)
    if p == "custom":
        if not custom_from or not custom_to:
            raise ValueError("custom preset requires from and to")
        return _utc(custom_from), _utc(custom_to)
    raise ValueError(f"Unknown preset '{preset}'. Valid: {', '.join(PRESETS)}")


def resolve_range(preset: str | None = None, from_val: str | None = None,
                  to_val: str | None = None, now: datetime | None = None,
                  fy_start_month: int = INDIA_FY_START_MONTH) -> tuple[datetime, datetime]:
    """Explicit from/to wins; otherwise resolve the named preset (#27/#56)."""
    if from_val and to_val:
        return _utc(from_val), _utc(to_val)
    if preset:
        return resolve_preset(preset, now=now, fy_start_month=fy_start_month,
                              custom_from=from_val, custom_to=to_val)
    raise ValueError("Provide from/to or a preset.")


def filtered_orders(db, business_id: str, start: datetime, end: datetime,
                    status: str | None = None, customer: str | None = None,
                    product: str | None = None, state: str | None = None,
                    payment_method: str | None = None,
                    courier: str | None = None) -> list:
    """Every report API supports from/to/status/courier/payment_method/customer/product/state (#56)."""
    from app.models.order import Order, OrderItem
    from app.models.payment import Payment
    from app.models.shipment import Shipment
    q = db.query(Order).filter_by(business_id=business_id).filter(
        Order.order_date >= _utc(start), Order.order_date < _utc(end))
    if status:
        s = status.upper()
        q = q.filter((Order.financial_status == s) | (Order.operational_status == s))
    if customer:
        q = q.filter(Order.customer_id == customer)
    if state:
        q = q.filter((Order.ship_state_code == state.upper()) |
                     (Order.place_of_supply == state.upper()))
    orders = q.all()
    if product:
        oids = {i.order_id for i in db.query(OrderItem).filter_by(
            business_id=business_id).filter(
            (OrderItem.product_id == product) | (OrderItem.sku == product)).all()}
        orders = [o for o in orders if o.id in oids]
    if payment_method:
        oids = {p.order_id for p in db.query(Payment).filter_by(
            business_id=business_id, method=payment_method).all()}
        orders = [o for o in orders if o.id in oids]
    if courier:
        oids = {s.order_id for s in db.query(Shipment).filter_by(
            business_id=business_id, carrier_code=courier).all()}
        orders = [o for o in orders if o.id in oids]
    return orders


# --- GST validation (#49/#50) ---

def gst_invoice_check(taxable, cgst, sgst, igst, total) -> dict:
    """Pure check: taxable + CGST + SGST + IGST == invoice total."""
    t, c, s, i, tot = (_d(v) for v in (taxable, cgst, sgst, igst, total))
    diff = (t + c + s + i) - tot
    return {"taxable": format(t, ".2f"), "cgst": format(c, ".2f"),
            "sgst": format(s, ".2f"), "igst": format(i, ".2f"),
            "total": format(tot, ".2f"), "diff": format(diff, ".2f"),
            "ok": abs(diff) <= Decimal("0.01")}


def validate_order_gst(db, business_id: str, order_id: str) -> dict:
    """Line-item GST validation with same/inter-state rule (#50).

    Same-state (place_of_supply == seller state): CGST+SGST, IGST must be 0.
    Inter-state: IGST only, CGST/SGST must be 0.
    Jurisdiction is advisory — final treatment needs CA review (flagged, never silent).
    """
    from app.models.order import Order, OrderItem
    o = db.query(Order).filter_by(business_id=business_id, id=order_id).first()
    if o is None:
        raise ValueError("Order not found.")
    items = db.query(OrderItem).filter_by(business_id=business_id, order_id=order_id).all()
    taxable = sum((_d(i.taxable_amount) for i in items), Decimal("0"))
    cgst = sum((_d(i.cgst_amount) for i in items), Decimal("0"))
    sgst = sum((_d(i.sgst_amount) for i in items), Decimal("0"))
    igst = sum((_d(i.igst_amount) for i in items), Decimal("0"))
    inv = gst_invoice_check(taxable, cgst, sgst, igst, _d(o.total_amount))
    seller = (o.business_state_code or "").upper() or None
    if seller is None:
        try:
            from app.models.business import Business
            b = db.query(Business).filter_by(id=business_id).first()
            seller = (getattr(b, "state_code", "") or "").upper() or None
        except Exception:
            seller = None
    buyer = ((o.place_of_supply or o.ship_state_code) or "").upper() or None
    errors: list[dict] = []
    warnings: list[dict] = []
    if not inv["ok"]:
        errors.append({"code": "INVOICE_TOTAL_MISMATCH",
                       "message": f"taxable+CGST+SGST+IGST != total (diff {inv['diff']})."})
    multi_rates = {str(i.gst_rate) for i in items if i.gst_rate is not None}
    if len(multi_rates) > 1:
        warnings.append({"code": "MULTI_RATE_ORDER",
                         "message": f"Multiple GST rates in one order: {sorted(multi_rates)}."})
    if any(i.hsn_code is None for i in items):
        warnings.append({"code": "HSN_MISSING",
                         "message": "HSN/SAC missing on one or more line items."})
    jurisdiction = "UNKNOWN"
    if seller and buyer:
        jurisdiction = "SAME_STATE" if seller == buyer else "INTER_STATE"
        if jurisdiction == "SAME_STATE" and igst > 0:
            errors.append({"code": "GST_JURISDICTION_MISMATCH",
                           "message": "Same-state sale must use CGST+SGST, not IGST."})
        if jurisdiction == "INTER_STATE" and (cgst > 0 or sgst > 0):
            errors.append({"code": "GST_JURISDICTION_MISMATCH",
                           "message": "Inter-state sale must use IGST, not CGST/SGST."})
    else:
        warnings.append({"code": "JURISDICTION_UNKNOWN",
                         "message": "Seller/buyer state unknown — intra-state split assumed; CA review required."})
    warnings.append({"code": "CA_REVIEW",
                     "message": "GST treatment must be reviewed by the business's CA/tax professional before filing."})
    return {"order_id": order_id, "invoice_check": inv,
            "jurisdiction": jurisdiction, "seller_state": seller, "buyer_state": buyer,
            "line_items": len(items), "valid": not errors,
            "errors": errors, "warnings": warnings,
            "display_timezone": DISPLAY_TZ}


def gst_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    """GST report over filtered orders: per-order invoice + jurisdiction checks."""
    orders = filtered_orders(db, business_id, start, end, **filters)
    rows, t_taxable, t_cgst, t_sgst, t_igst = [], Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0")
    invalid = 0
    for o in orders:
        try:
            v = validate_order_gst(db, business_id, o.id)
        except ValueError:
            continue
        if not v["valid"]:
            invalid += 1
        t_taxable += _d(v["invoice_check"]["taxable"])
        t_cgst += _d(v["invoice_check"]["cgst"])
        t_sgst += _d(v["invoice_check"]["sgst"])
        t_igst += _d(v["invoice_check"]["igst"])
        rows.append({"order_id": o.id, "order_name": o.shopify_order_name,
                     "valid": v["valid"], "jurisdiction": v["jurisdiction"],
                     "invoice_check": v["invoice_check"],
                     "errors": v["errors"], "warnings": [w["code"] for w in v["warnings"]]})
    return {"period": {"from": _utc(start).isoformat(), "to": _utc(end).isoformat()},
            "orders": len(orders), "invalid": invalid,
            "totals": {"taxable": format(t_taxable, ".2f"), "cgst": format(t_cgst, ".2f"),
                       "sgst": format(t_sgst, ".2f"), "igst": format(t_igst, ".2f")},
            "rows": rows, "display_timezone": DISPLAY_TZ}


# --- Ledger backfill (owed from Task 1) ---

def backfill_ledger(db, business_id: str) -> dict:
    """Create ledger events for pre-existing orders/payments/refunds.

    Idempotent: reuses the same idempotency keys as the live hooks, so a
    second run creates nothing new. Returns per-type created/skipped counts.
    """
    from app.models.order import Order
    from app.models.payment import Payment
    from app.models.refund import Refund
    from app.models.financial_transaction import FinancialTransaction
    from app.services import ledger_service as _ls
    out = {"sales": {"created": 0, "skipped": 0},
           "payments": {"created": 0, "skipped": 0},
           "refunds": {"created": 0, "skipped": 0}}
    existing = {r.idempotency_key for r in db.query(FinancialTransaction).filter_by(
        business_id=business_id).all()}
    try:
        paid_oids = {p.order_id for p in db.query(Payment).filter_by(
            business_id=business_id).all()
            if (p.payment_status or "").upper() in ("PAID", "COMPLETED", "SETTLED", "SUCCESS")}
    except Exception:
        paid_oids = set()
    for o in db.query(Order).filter_by(business_id=business_id).all():
        if o.cancelled_at is not None and o.id not in paid_oids:
            # #43/#100: cancelled before payment -> no revenue, no receipt, no refund.
            out["sales"]["skipped"] += 1
            continue
        if f"SALE:{o.id}" in existing:
            out["sales"]["skipped"] += 1
            continue
        try:
            _ls.record_sale_from_order(db, business_id, o)
            existing.add(f"SALE:{o.id}")
            out["sales"]["created"] += 1
        except Exception:
            out["sales"]["skipped"] += 1
    for p in db.query(Payment).filter_by(business_id=business_id).all():
        if f"PAYMENT:{p.id}" in existing:
            out["payments"]["skipped"] += 1
            continue
        try:
            _ls.record_payment_from_payment(db, business_id, p)
            existing.add(f"PAYMENT:{p.id}")
            out["payments"]["created"] += 1
        except Exception:
            out["payments"]["skipped"] += 1
    for r in db.query(Refund).filter_by(business_id=business_id).all():
        if f"REFUND:{r.id}" in existing:
            out["refunds"]["skipped"] += 1
            continue
        try:
            _ls.record_refund_from_refund(db, business_id, r)
            existing.add(f"REFUND:{r.id}")
            out["refunds"]["created"] += 1
        except Exception:
            out["refunds"]["skipped"] += 1
    db.commit()
    out["total_created"] = out["sales"]["created"] + out["payments"]["created"] + out["refunds"]["created"]
    return out


# --- Unified report builders wired to the finance engines (#55) ---

def _period_engines(db, business_id: str, start: datetime, end: datetime,
                    order_ids=None) -> dict:
    from app.services import ledger_service as _ls
    return _ls.period_summary(db, business_id, _utc(start), _utc(end), order_ids=order_ids)


def dashboard_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    """KPI cards per #28, wired to ledger/bank/tally engines (no duplicated math)."""
    from app.models.shipment import Shipment
    from app.models.refund import Refund
    from app.models.payment import Payment
    orders = filtered_orders(db, business_id, start, end, **filters)
    oids = {o.id for o in orders}
    eng = _period_engines(db, business_id, start, end, order_ids=oids)
    rev, profit = eng["revenue"], eng["profit"]
    refunds = db.query(Refund).filter_by(business_id=business_id).all()
    refunds = [r for r in refunds if r.order_id in oids]
    payments = db.query(Payment).filter_by(business_id=business_id).all()
    payments = [p for p in payments if p.order_id in oids]
    paid = sum(float(p.amount or 0) for p in payments
               if (p.payment_status or "").upper() in ("PAID", "COMPLETED", "SETTLED", "SUCCESS"))
    pending = sum(float(o.total_amount or 0) for o in orders) - paid
    ships = db.query(Shipment).filter_by(business_id=business_id).all()
    ships = [s for s in ships if s.order_id in oids]
    delivered = sum(1 for s in ships if (s.tracking_status or "") == "DELIVERED")
    in_transit = sum(1 for s in ships if (s.tracking_status or "") in (
        "IN_TRANSIT", "AT_HUB", "OUT_FOR_DELIVERY", "BOOKED"))
    rto = sum(1 for s in ships if "RTO" in (s.tracking_status or ""))
    margin = profit.get("margin_pct", "0.00")
    return {
        "period": {"from": _utc(start).isoformat(), "to": _utc(end).isoformat()},
        "kpis": {
            "orders_total": len(orders),
            "successful_orders": sum(1 for o in orders if (o.financial_status or "") == "PAID"),
            "cancelled_orders": sum(1 for o in orders if o.cancelled_at is not None),
            "refunded_orders": len({r.order_id for r in refunds}),
            "gross_sales": rev["gross_inclusive"],
            "discounts": rev["discounts"],
            "taxable_sales": rev["gross_exclusive"],
            "gst": rev["gst"],
            "net_sales": rev["net_exclusive"],
            "payments_received": format(_d(paid), ".2f"),
            "payments_pending": format(_d(pending), ".2f"),
            "cogs": profit["cogs"],
            "shipping_cost": profit["shipping"],
            "gateway_fees": profit["gateway_fees"],
            "other_expenses": profit["other"],
            "gross_profit": profit["gross_profit"],
            "operating_profit": profit["operating_profit"],
            "profit_label": profit["label"],
            "profit_warning": profit["warning"],
            "profit_margin_pct": margin,
            "delivered_shipments": delivered,
            "in_transit": in_transit,
            "rto": rto,
        },
        "display_timezone": DISPLAY_TZ,
    }


def sales_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    orders = filtered_orders(db, business_id, start, end, **filters)
    eng = _period_engines(db, business_id, start, end, order_ids={o.id for o in orders})
    by_status: dict[str, int] = {}
    for o in orders:
        by_status[o.financial_status or "UNKNOWN"] = by_status.get(o.financial_status or "UNKNOWN", 0) + 1
    rows = [{"order_id": o.id, "name": o.shopify_order_name,
             "date": o.order_date.isoformat() if o.order_date else None,
             "financial": o.financial_status, "operational": o.operational_status,
             "total": str(o.total_amount or 0)} for o in orders]
    return {"period": {"from": _utc(start).isoformat(), "to": _utc(end).isoformat()},
            "orders": len(orders), "by_status": by_status,
            "revenue": eng["revenue"], "rows": rows, "display_timezone": DISPLAY_TZ}


def orders_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    return sales_report(db, business_id, start, end, **filters)


def payments_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    from app.models.payment import Payment
    orders = filtered_orders(db, business_id, start, end, **filters)
    oids = {o.id for o in orders}
    pays = [p for p in db.query(Payment).filter_by(business_id=business_id).all() if p.order_id in oids]
    by_status: dict[str, int] = {}
    for p in pays:
        by_status[p.payment_status or "UNKNOWN"] = by_status.get(p.payment_status or "UNKNOWN", 0) + 1
    return {"period": {"from": _utc(start).isoformat(), "to": _utc(end).isoformat()},
            "payments": len(pays), "by_status": by_status,
            "total_received": format(sum((_d(p.amount) for p in pays
                                          if (p.payment_status or '').upper() not in ('FAILED', 'CANCELLED')),
                                         Decimal("0")), ".2f"),
            "rows": [{"id": p.id, "order_id": p.order_id, "amount": str(p.amount),
                      "status": p.payment_status, "method": p.method} for p in pays],
            "display_timezone": DISPLAY_TZ}


def refunds_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    from app.models.refund import Refund
    orders = filtered_orders(db, business_id, start, end, **filters)
    oids = {o.id for o in orders}
    refs = [r for r in db.query(Refund).filter_by(business_id=business_id).all() if r.order_id in oids]
    return {"period": {"from": _utc(start).isoformat(), "to": _utc(end).isoformat()},
            "refunds": len(refs),
            "total": format(sum((_d(r.amount) for r in refs), Decimal("0")), ".2f"),
            "rows": [{"id": r.id, "order_id": r.order_id, "amount": str(r.amount),
                      "status": r.status} for r in refs],
            "display_timezone": DISPLAY_TZ}


def profit_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    orders = filtered_orders(db, business_id, start, end, **filters)
    eng = _period_engines(db, business_id, start, end, order_ids={o.id for o in orders})
    return {"period": {"from": _utc(start).isoformat(), "to": _utc(end).isoformat()},
            "orders": len(orders), "revenue": eng["revenue"], "profit": eng["profit"],
            "cogs_total": eng["cogs_total"], "transaction_count": eng["transaction_count"],
            "display_timezone": DISPLAY_TZ}


def courier_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    from app.models.shipment import Shipment
    from app.services.shipment_service import is_awaiting_awb
    orders = filtered_orders(db, business_id, start, end, **filters)
    oids = {o.id for o in orders}
    # Excluded rather than masked. This report counts parcels handed to a
    # courier and breaks the total down per carrier; an awaiting shipment has
    # reached neither, so keeping it - even with its AWB blanked - would
    # inflate "shipments" and add a phantom row under India Post. A separate
    # "waiting for a tracking number" count would be the honest presentation,
    # but that is a report shape change this task does not own.
    ships = [s for s in db.query(Shipment).filter_by(business_id=business_id).all()
             if s.order_id in oids and not is_awaiting_awb(s.awb_number)]
    by_courier: dict[str, dict] = {}
    for s in ships:
        c = by_courier.setdefault(s.carrier_code, {"shipments": 0, "delivered": 0,
                                                    "in_transit": 0, "rto": 0})
        c["shipments"] += 1
        st = s.tracking_status or ""
        if st == "DELIVERED":
            c["delivered"] += 1
        elif st in ("IN_TRANSIT", "AT_HUB", "OUT_FOR_DELIVERY", "BOOKED"):
            c["in_transit"] += 1
        if "RTO" in st:
            c["rto"] += 1
    return {"period": {"from": _utc(start).isoformat(), "to": _utc(end).isoformat()},
            "shipments": len(ships), "by_courier": by_courier,
            "rows": [{"awb": s.awb_number, "carrier": s.carrier_code,
                      "status": s.tracking_status, "order_id": s.order_id} for s in ships],
            "display_timezone": DISPLAY_TZ}


def reconciliation_report(db, business_id: str, start: datetime, end: datetime, **filters) -> dict:
    from app.models.reconciliation import Reconciliation
    orders = filtered_orders(db, business_id, start, end, **filters)
    oids = {o.id for o in orders}
    exc = [r for r in db.query(Reconciliation).filter_by(business_id=business_id).all()
           if r.order_id in oids]
    open_exc = [r for r in exc if not r.resolved]
    by_sev: dict[str, int] = {}
    for r in open_exc:
        by_sev[r.severity] = by_sev.get(r.severity, 0) + 1
    matched = len(oids) - len({r.order_id for r in open_exc if r.order_id})
    return {"period": {"from": _utc(start).isoformat(), "to": _utc(end).isoformat()},
            "orders": len(orders), "matched": matched,
            "open": len(open_exc), "by_severity": by_sev,
            "rows": [{"order_id": r.order_id, "code": r.code, "severity": r.severity,
                      "message": r.message, "resolved": r.resolved} for r in exc],
            "display_timezone": DISPLAY_TZ}


ASYNC_EXPORT_ROW_LIMIT = 5000


def export_table_csv(headers: list[str], rows: list[list]) -> bytes:
    import csv as _csv
    import io as _io
    buf = _io.StringIO()
    w = _csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8")


def _num(v):
    """Coerce money cells to real numerics (#79): '1180.00' -> 1180.0, else as-is."""
    if isinstance(v, bool) or v is None or isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", "").strip())
    except (ValueError, TypeError):
        return v


def _xldate(v):
    """Coerce ISO strings to naive-IST datetimes so Excel holds real dates (#79).

    Repo display discipline is Asia/Kolkata: convert UTC->IST first.
    """
    from datetime import timedelta
    if v is None:
        return v
    if isinstance(v, datetime):
        d = v if v.tzinfo else v.replace(tzinfo=timezone.utc)
    else:
        try:
            d = datetime.fromisoformat(str(v))
            d = d if d.tzinfo else d.replace(tzinfo=timezone.utc)
        except (ValueError, TypeError):
            return v
    d = d.astimezone(timezone.utc)
    try:
        from zoneinfo import ZoneInfo
        d = d.astimezone(ZoneInfo("Asia/Kolkata"))
    except Exception:
        d = d + timedelta(hours=5, minutes=30)
    return d.replace(tzinfo=None)


def export_table_xlsx(title: str, headers: list[str], rows: list[list],
                      money_cols: tuple = (), date_cols: tuple = ()) -> bytes:
    import io as _io
    from openpyxl import Workbook as _WB
    from openpyxl.styles import Font as _Font, PatternFill as _Fill
    wb = _WB()
    ws = wb.active
    ws.title = title[:31]
    ws.freeze_panes = "A2"
    ws.append(headers)
    for c in ws[1]:
        c.font = _Font(bold=True, color="FFFFFF")
        c.fill = _Fill("solid", fgColor="1F4E5F")
    ws.auto_filter.ref = ws.dimensions
    for r in rows:
        out = list(r)
        for ci in money_cols:
            if 0 < ci <= len(out):
                out[ci - 1] = _num(out[ci - 1])
        for ci in date_cols:
            if 0 < ci <= len(out):
                out[ci - 1] = _xldate(out[ci - 1])
        ws.append(out)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                if c.column in money_cols or not money_cols:
                    c.number_format = "#,##0.00"
            if isinstance(c.value, datetime) and c.column in date_cols:
                c.number_format = "YYYY-MM-DD"
    if rows and money_cols:  # totals row with live SUM formulas (#79)
        n = ws.max_row
        ws.cell(row=n + 1, column=1, value="TOTAL").font = _Font(bold=True)
        for ci in money_cols:
            letter = ws.cell(row=1, column=ci).column_letter
            cell = ws.cell(row=n + 1, column=ci, value=f"=SUM({letter}2:{letter}{n})")
            cell.font = _Font(bold=True)
            cell.number_format = "#,##0.00"
    for col in ws.columns:
        w = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(w + 2, 40)
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def month_window(month: str) -> tuple[datetime, datetime]:
    y, m = [int(x) for x in month.split("-")]
    start = datetime(y, m, 1, tzinfo=timezone.utc)
    if m == 12:
        end = datetime(y + 1, 1, 1, tzinfo=timezone.utc)
    else:
        end = datetime(y, m + 1, 1, tzinfo=timezone.utc)
    return start, end


def monthly_report(db, business_id: str, month: str) -> dict:
    from sqlalchemy import func
    from app.models.order import Order, OrderItem
    from app.models.shipment import Shipment
    from app.models.return_record import ReturnRecord
    from app.models.refund import Refund
    from app.models.reconciliation import Reconciliation
    from app.models.sla import ShipmentFinancial
    from app.services.cost_service import get_cost
    from app.services.shipment_service import is_awaiting_awb
    start, end = month_window(month)
    from datetime import timedelta
    month_end = end - timedelta(microseconds=1)
    orders = db.query(Order).filter_by(business_id=business_id).filter(
        Order.order_date >= start, Order.order_date < end).all()
    oids = [o.id for o in orders]

    def cnt(pred):
        return sum(1 for o in orders if pred(o))

    ops = [(o.operational_status or "") for o in orders]
    order_sec = {"total": len(orders),
                 "paid": cnt(lambda o: (o.financial_status or "") == "PAID"),
                 "unpaid": cnt(lambda o: (o.financial_status or "") not in ("PAID", "REFUNDED")),
                 "cancelled": cnt(lambda o: o.cancelled_at is not None),
                 "dispatched": sum(ops.count(s) for s in ("DISPATCHED",)),
                 "delivered": 0, "returned": 0, "rto": 0}
    ships = db.query(Shipment).filter_by(business_id=business_id).all()
    # Awaiting shipments are excluded, matching courier_report above and the
    # shipment_rows filter below. Every order synced from Shopify now owns one, so
    # leaving them in invented a courier row for a carrier the parcel has not
    # reached and inflated that courier's shipment and breached counts.
    in_month = [s for s in ships
                if s.order_id in set(oids) and not is_awaiting_awb(s.awb_number)]
    courier_sec: dict[str, dict] = {}
    for s in in_month:
        c = courier_sec.setdefault(s.carrier_code, {"shipments": 0, "delivered": 0, "in_transit": 0,
                                                    "rto": 0, "returned": 0, "breached": 0})
        c["shipments"] += 1
        st = s.tracking_status or ""
        if st == "DELIVERED":
            c["delivered"] += 1
            order_sec["delivered"] += 1
        elif st in ("IN_TRANSIT", "AT_HUB", "OUT_FOR_DELIVERY"):
            c["in_transit"] += 1
        if "RTO" in st:
            c["rto"] += 1
            order_sec["rto"] += 1
        if st in ("RETURNED", "RETURN_AT_HUB"):
            c["returned"] += 1
            order_sec["returned"] += 1
    from app.services.sla_service import sla_status, shipment_clock
    now = datetime.now(timezone.utc)
    for s in in_month:
        if sla_status(shipment_clock(db, s), now)["status"] == "BREACHED":
            courier_sec[s.carrier_code]["breached"] += 1

    rets = db.query(ReturnRecord).filter_by(business_id=business_id).filter(
        ReturnRecord.order_id.in_(oids)).all() if oids else []
    return_sec = {"total": len(rets),
                  "customer": sum(1 for r in rets if r.return_type == "CUSTOMER_RETURN"),
                  "rto": sum(1 for r in rets if r.return_type == "RTO"),
                  "partial": sum(1 for r in rets if r.return_type == "PARTIAL_RETURN"),
                  "refunded": 0, "pending_refund": 0}
    refunded_oids = {r.order_id for r in db.query(Refund).filter_by(business_id=business_id).all()}
    for r in rets:
        if r.order_id in refunded_oids:
            return_sec["refunded"] += 1
        else:
            return_sec["pending_refund"] += 1

    gross = _d(sum((_d(o.total_amount) for o in orders), _d(0)))
    discounts = _d(sum((_d(o.discount_amount) for o in orders), _d(0)))
    refunds = _d(sum((_d(r.amount) for r in db.query(Refund).filter_by(business_id=business_id).all()
                        if r.order_id in set(oids)), _d(0)))
    fin_rows = db.query(ShipmentFinancial).filter_by(business_id=business_id).all()
    fin_rows = [f for f in fin_rows if f.order_id in set(oids)]
    expected = _d(sum((_d(f.expected_cod_amount) for f in fin_rows), _d(0)))
    collected = _d(sum((_d(f.collected_amount) for f in fin_rows), _d(0)))
    settled = _d(sum((_d(f.settled_amount) for f in fin_rows), _d(0)))
    fees = _d(sum((_d(f.fee_amount) for f in fin_rows), _d(0)))
    money_sec = {"gross": float(gross), "discounts": float(discounts), "refunds": float(refunds),
                 "expected": float(expected), "collected": float(collected),
                 "settled": float(settled), "pending": float(expected - settled),
                 "fees": float(fees),
                 "net": float(settled - fees)}

    exc = db.query(Reconciliation).filter_by(business_id=business_id, resolved=False).all()
    exc = [r for r in exc if r.order_id in set(oids)]
    by_sev: dict[str, int] = {}
    for r in exc:
        by_sev[r.severity] = by_sev.get(r.severity, 0) + 1
    exc_sec = {"open": len(exc), "by_severity": by_sev}

    cost_lines = []
    cost_total = _d(0)
    order_count = max(len(orders), 1)
    item_qty = sum(i.quantity for i in db.query(OrderItem).filter(OrderItem.order_id.in_(oids)).all()) if oids else 0
    multipliers = {"COGS_DEFAULT": item_qty, "SHIPPING": len(orders), "GATEWAY_FEE": len(orders),
                   "PACKAGING": len(orders), "RETURN_COST": return_sec["total"],
                   "RTO_COST": return_sec["rto"], "OTHER": 1}
    sources = set()
    for key in ("COGS_DEFAULT", "SHIPPING", "GATEWAY_FEE", "PACKAGING", "RETURN_COST", "RTO_COST", "OTHER"):
        amt, src = get_cost(db, business_id, key, month_end)
        line = _d(_d(amt) * multipliers[key])
        cost_total += line
        sources.add(src)
        cost_lines.append({"key": key, "per_unit": float(_d(amt)), "units": multipliers[key],
                           "total": float(line), "source": src})
    net_revenue = gross - discounts - refunds
    profit = net_revenue - cost_total
    label = "OPERATING PROFIT" if sources == {"ACTUAL_STATEMENT"} else "ESTIMATED OPERATING PROFIT"
    pnl = {"gross": float(gross), "discounts": float(discounts), "refunds": float(refunds),
           "net_revenue": float(net_revenue),
           "costs": cost_lines, "cost_total": float(cost_total), "profit": float(profit), "label": label}

    return {"month": month, "generated_at": now.isoformat(), "orders": order_sec, "courier": courier_sec,
            "returns": return_sec, "money": money_sec, "exceptions": exc_sec,
            "costs": {"lines": cost_lines, "total": float(cost_total)}, "profitability": pnl,
            "order_rows": [{"name": o.shopify_order_name, "financial": o.financial_status,
                            "operational": o.operational_status, "total": float(_d(o.total_amount))} for o in orders],
            "shipment_rows": [{"awb": s.awb_number, "carrier": s.carrier_code, "status": s.tracking_status,
                               "location": s.current_location or ""}
                              for s in ships if not is_awaiting_awb(s.awb_number)]}
