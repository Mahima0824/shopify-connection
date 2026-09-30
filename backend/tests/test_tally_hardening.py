"""Task 3 TDD (RED first): Tally export hardening — #82 matrix + duplicate + validation gate."""
from datetime import datetime, timezone


def _mkbiz():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from app.database import Base
    import app.models  # noqa: F401
    from app.models.business import Business
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(eng)
    mk = sessionmaker(bind=eng)
    db = mk()
    b = Business(name="B3", email="b3@t.in")
    db.add(b)
    db.commit()
    db.refresh(b)
    return db, b


_UID = {"n": 0}


def _mkorder(db, b, name="#T1", total="1180.00", tax="180.00", customer_id=None,
             paid=True, method="ONLINE", dt=None):
    from app.models.order import Order
    from app.models.customer import Customer
    from app.models.payment import Payment
    from app.services import ledger_service as ls
    _UID["n"] += 1
    uniq = f"{name}-{_UID['n']}"
    dt = dt or datetime(2026, 9, 10, 10, 0, tzinfo=timezone.utc)
    c = None
    if customer_id != "MISSING":
        c = Customer(business_id=b.id, first_name="C", last_name="T", email="c@t.in")
        db.add(c)
        db.commit()
        db.refresh(c)
    o = Order(business_id=b.id, internal_order_number=f"ORD-{uniq}",
              shopify_order_id=f"gid://{uniq}", shopify_order_name=name,
              customer_id=c.id if c else None,
              currency="INR", subtotal_amount=float(total) - float(tax),
              discount_amount=0, shipping_amount=0, tax_amount=float(tax),
              total_amount=float(total), payment_status="PAID" if paid else "PENDING",
              financial_status="PAID" if paid else "PENDING",
              fulfillment_status="UNFULFILLED", operational_status="NEW", order_date=dt)
    db.add(o)
    db.commit()
    db.refresh(o)
    sale = ls.record_event(db, b.id, "SALE", total, tax, transaction_date=dt,
                           order_id=o.id, reference_number=name,
                           idempotency_key=f"SALE:{o.id}", tally_voucher_type="Sales",
                           tally_voucher_number=name)
    db.commit()
    pay = None
    if paid:
        pay = Payment(business_id=b.id, order_id=o.id, amount=float(total),
                      payment_status="PAID", method=method, transaction_id=f"pay_{name}")
        db.add(pay)
        db.commit()
        db.refresh(pay)
        ls.record_event(db, b.id, "PAYMENT", total, "0", transaction_date=dt,
                        order_id=o.id, payment_id=pay.id, payment_method=method,
                        reference_number=pay.transaction_id,
                        idempotency_key=f"PAYMENT:{pay.id}", tally_voucher_type="Receipt")
        db.commit()
    return o, (c.id if c else None)


def test_validation_gate_blocks_on_errors():
    from app.services import tally_service as ts
    db, b = _mkbiz()
    _mkorder(db, b, name="#DUP", total="1180.00", tax="180.00")
    # second order with same invoice name -> duplicate invoice error
    _mkorder(db, b, name="#DUP", total="500.00", tax="0.00")
    res = ts.validate_export(db, b.id)
    assert res["errors"], "expected blocking errors for duplicate invoice"
    assert res["transactions"] >= 2
    try:
        ts.generate_workbook_export(db, b.id, "u1")
        assert False, "export must be blocked while errors exist"
    except Exception as e:
        assert "VALIDATION" in str(e) or "validation" in str(e).lower() or getattr(e, "code", "") != ""


def test_cancelled_unpaid_excluded():
    from app.services import tally_service as ts
    from app.services import ledger_service as ls
    db, b = _mkbiz()
    o, _ = _mkorder(db, b, name="#CX", total="900.00", tax="0.00", paid=False)
    ls.record_cancellation(db, b.id, o, refund_required=False)
    db.commit()
    rows = ts.collect_export_rows(db, b.id)
    kinds = {(r["voucher_type"], r["transaction_type"]) for r in rows}
    assert kinds == set() or all(r["order_ref"] != "#CX" or r["amount"] == 0 for r in rows), kinds
    sales = [r for r in rows if r["voucher_type"] == "Sales"]
    assert not [r for r in sales if r["order_ref"] == "#CX"], "cancelled-unpaid must produce no sale"


def test_workbook_multisheet_numeric_dates():
    from app.services import tally_service as ts
    from app.services import ledger_service as ls
    from datetime import datetime as _dt
    db, b = _mkbiz()
    o, _ = _mkorder(db, b, name="#W1", total="1180.00", tax="180.00")
    ls.record_event(db, b.id, "PAYMENT_GATEWAY_FEE", "20.00", "0",
                    transaction_date=_dt(2026, 9, 10, 12, 0, tzinfo=timezone.utc),
                    order_id=o.id, idempotency_key=f"FEE:{o.id}")
    ls.record_event(db, b.id, "SHIPPING_EXPENSE", "70.00", "0",
                    transaction_date=_dt(2026, 9, 11, 12, 0, tzinfo=timezone.utc),
                    order_id=o.id, idempotency_key=f"SHIP:{o.id}")
    db.commit()
    out = ts.generate_workbook_export(db, b.id, "u1")
    assert out["file_name"].startswith("TALLY_EXPORT_2026_09")
    from openpyxl import load_workbook
    import io
    wb = load_workbook(filename=io.BytesIO(out["content"]))
    for sheet in ("Summary", "Sales", "Sales_Items", "Receipts", "Credit_Notes", "Expenses", "Journal"):
        assert sheet in wb.sheetnames, wb.sheetnames
    ws = wb["Sales"]
    assert ws.freeze_panes is not None and ws.freeze_panes != "A1"
    assert ws.auto_filter is not None
    # numeric amount cell must be a number, date cell a date
    found_num = found_date = False
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, (int, float)) and c.value > 0:
                found_num = True
                assert not isinstance(c.value, str)
        dv = row[2].value if len(row) > 2 else None
        from datetime import datetime as _DT
        if isinstance(dv, _DT):
            found_date = True
    assert found_num, "numeric cells must be numbers"
    assert found_date, "date cells must be Excel dates"


def test_duplicate_export_blocked():
    from app.services import tally_service as ts
    db, b = _mkbiz()
    _mkorder(db, b, name="#U1", total="500.00", tax="0.00")
    ts.generate_workbook_export(db, b.id, "u1")
    db.commit()
    try:
        ts.generate_workbook_export(db, b.id, "u1")
        assert False, "second export of same rows must be blocked"
    except Exception as e:
        assert getattr(e, "code", "") == "DUPLICATE_EXPORT" or "duplicate" in str(e).lower()


def test_batch_lifecycle():
    from app.services import tally_service as ts
    db, b = _mkbiz()
    _mkorder(db, b, name="#B1", total="500.00", tax="0.00")
    out = ts.generate_workbook_export(db, b.id, "u1")
    assert out["batch"]["status"] == "GENERATED"
    b1 = ts.mark_downloaded(db, b.id, out["batch"]["id"])
    assert b1["status"] == "DOWNLOADED"
    b2 = ts.mark_imported(db, b.id, out["batch"]["id"], imported=True)
    assert b2["status"] == "IMPORTED"
    # partial path
    _mkorder(db, b, name="#B2", total="600.00", tax="0.00")
    out2 = ts.generate_workbook_export(db, b.id, "u1")
    b3 = ts.mark_imported(db, b.id, out2["batch"]["id"], imported=False, partial=True)
    assert b3["status"] == "PARTIALLY_IMPORTED"


def test_api_workbook_lifecycle_envelope_and_from_alias():
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from app.database import Base, get_db
    import app.models  # noqa: F401
    from app.main import app
    from app.models.business import Business
    from app.models.user import User
    from app.services.auth_service import hash_password
    from app.services import ledger_service as ls
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(eng)
    mk = sessionmaker(bind=eng)
    db = mk()
    b = Business(name="BAPI", email="bapi@t.in")
    db.add(b)
    db.commit()
    db.refresh(b)
    u = User(business_id=b.id, name="A", email="a3@t.in",
             password_hash=hash_password("x"), role="ACCOUNTANT")
    db.add(u)
    db.commit()
    bid = b.id
    db.close()
    app.dependency_overrides[get_db] = lambda: mk()
    c = TestClient(app)
    tok = c.post("/api/v1/auth/login", json={"email": "a3@t.in", "password": "x"}).json()["data"]["token"]
    h = {"Authorization": f"Bearer {tok}"}

    # seed one sale via API-scoped db
    from types import SimpleNamespace
    db2 = mk()
    _mkorder(db2, SimpleNamespace(id=bid), name="#API1", total="1180.00", tax="180.00")
    db2.commit()
    db2.close()

    # validate with `from` alias per #56 contract
    r = c.post("/api/v1/tally/validate?from=2026-09-01T00:00:00%2B00:00", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["success"] is True
    assert r.json()["data"]["can_export"] is True

    # workbook download
    r = c.post("/api/v1/tally/export-workbook", headers=h)
    assert r.status_code == 200, r.text
    assert "TALLY_EXPORT_" in r.headers.get("content-disposition", "")
    assert r.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats-officedocument")

    # lifecycle: downloaded -> imported
    items = c.get("/api/v1/tally/exports", headers=h).json()["data"]["items"]
    assert len(items) >= 1
    bid = items[0]["id"]
    assert c.post(f"/api/v1/tally/exports/{bid}/downloaded", headers=h).json()["data"]["status"] == "DOWNLOADED"
    r = c.post(f"/api/v1/tally/exports/{bid}/mark-imported",
               json={"imported": True}, headers=h)
    assert r.json()["data"]["status"] == "IMPORTED"

    # re-export now blocked with envelope error (duplicate prevention #44)
    r = c.post("/api/v1/tally/export-workbook", headers=h)
    assert r.status_code == 409
    body = r.json()
    assert body["success"] is False
    assert body["error"]["code"] == "DUPLICATE_EXPORT"
    assert body["error"]["message"]

    # unknown batch -> envelope NOT_FOUND
    r = c.post("/api/v1/tally/exports/nope/mark-imported",
               json={"imported": True}, headers=h)
    assert r.status_code == 404
    assert r.json() == {"success": False, "error": {"code": "NOT_FOUND",
                                                   "message": "Export batch not found."}}


def test_full_matrix_82_sale_cancel_refund_fee_shipping_cod_online():
    """#82 matrix: sale, cancelled-unpaid, cancelled-paid, full + partial refund,
    gateway fee, shipping, COD, online (inter/intra-state documented as warning)."""
    from datetime import datetime as _dt
    from app.services import tally_service as ts
    from app.services import ledger_service as ls
    from app.models.refund import Refund
    db, b = _mkbiz()

    # 1. successful online same-state sale
    o1, _ = _mkorder(db, b, name="#M1", total="1180.00", tax="180.00", method="ONLINE")
    # 2. cancelled unpaid -> nothing exportable
    o2, _ = _mkorder(db, b, name="#M2", total="900.00", tax="0.00", paid=False)
    ls.record_cancellation(db, b.id, o2, refund_required=False)
    # 3. cancelled paid -> sale + payment + refund/credit
    o3, _ = _mkorder(db, b, name="#M3", total="1180.00", tax="180.00", method="ONLINE")
    r3 = Refund(business_id=b.id, order_id=o3.id, shopify_refund_id="rf-m3", amount=1180.0)
    db.add(r3)
    db.commit()
    ls.record_event(db, b.id, "REFUND", "1180.00", "180.00", order_id=o3.id,
                    refund_id=r3.id, idempotency_key="REFUND:rf-m3",
                    tally_voucher_type="Credit Note")
    # 4. full refund on o1's order would double-count; use partial refund instead
    r1p = Refund(business_id=b.id, order_id=o1.id, shopify_refund_id="rf-m1p", amount=590.0)
    db.add(r1p)
    db.commit()
    ls.record_event(db, b.id, "REFUND", "590.00", "90.00", order_id=o1.id,
                    refund_id=r1p.id, idempotency_key="REFUND:rf-m1p",
                    tally_voucher_type="Credit Note")
    # 5. gateway fee + 6. shipping on o1
    base = _dt(2026, 9, 12, 12, 0, tzinfo=timezone.utc)
    ls.record_event(db, b.id, "PAYMENT_GATEWAY_FEE", "20.00", "0",
                    transaction_date=base, order_id=o1.id, idempotency_key="FEE:m1")
    ls.record_event(db, b.id, "SHIPPING_EXPENSE", "70.00", "0",
                    transaction_date=base, order_id=o1.id, idempotency_key="SHIP:m1")
    # 7. COD sale
    _mkorder(db, b, name="#M4", total="590.00", tax="90.00", method="COD")
    db.commit()

    v = ts.validate_export(db, b.id)
    assert v["can_export"] is True, v["errors"]
    # inter/intra-state: place-of-supply not captured -> advisory warning present
    assert any(w["code"] == "STATE_UNVERIFIED" for w in v["warnings"])

    rows = ts.collect_export_rows(db, b.id)
    by_order = {}
    for r in rows:
        by_order.setdefault(r["order_ref"], set()).add(r["voucher_type"])
    assert "#M2" not in by_order, "cancelled-unpaid exports nothing"
    assert by_order["#M1"] >= {"Sales", "Receipt", "Credit Note", "Payment"}
    assert by_order["#M3"] >= {"Sales", "Receipt", "Credit Note"}
    assert by_order["#M4"] >= {"Sales", "Receipt"}
    cod_receipts = [r for r in rows if r["order_ref"] == "#M4" and r["voucher_type"] == "Receipt"]
    assert cod_receipts and cod_receipts[0]["payment_method"] == "COD"

    out = ts.generate_workbook_export(db, b.id, "u1")
    assert out["batch"]["status"] == "GENERATED"
    assert out["file_name"].startswith("TALLY_EXPORT_")


def test_mapping_completeness_and_separate_voucher_types():
    from app.services import tally_service as ts
    db, b = _mkbiz()
    comp = ts.ledger_mapping_completeness(db, b.id)
    assert comp["complete"] is True
    assert "sales" in comp["mappings"]
    o, _ = _mkorder(db, b, name="#V1", total="1180.00", tax="180.00")
    from app.models.refund import Refund
    r = Refund(business_id=b.id, order_id=o.id, shopify_refund_id="rf1", amount=590.0)
    db.add(r)
    db.commit()
    from app.services import ledger_service as ls
    ls.record_event(db, b.id, "REFUND", "590.00", "90.00", order_id=o.id,
                    refund_id=r.id, idempotency_key=f"REFUND:{r.id}",
                    tally_voucher_type="Credit Note")
    db.commit()
    rows = ts.collect_export_rows(db, b.id)
    vtypes = {x["voucher_type"] for x in rows}
    assert "Sales" in vtypes and "Receipt" in vtypes and "Credit Note" in vtypes
    assert "Generic" not in vtypes and "Voucher" not in vtypes
