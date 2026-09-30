"""Final fix-wave tests (reviewer MUST-FIX F1-F5)."""
from datetime import datetime, timezone


def _mkdb():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from app.database import Base
    import app.models  # noqa: F401
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    Base.metadata.create_all(eng)
    mk = sessionmaker(bind=eng)
    db = mk()
    from app.models.business import Business
    b = Business(name="FF", email="ff@t.in")
    db.add(b)
    db.commit()
    db.refresh(b)
    return mk, db, b


_UID = {"n": 0}


def _mksale(db, b, name, total="1180.00", tax="180.00", seller_state=None,
            buyer_state=None, cust_state=None, dt=None):
    """Order + SALE ledger event with jurisdiction fields."""
    from app.models.order import Order
    from app.models.customer import Customer
    from app.services import ledger_service as ls
    _UID["n"] += 1
    uniq = f"{name}-{_UID['n']}"
    dt = dt or datetime(2026, 9, 10, 10, 0, tzinfo=timezone.utc)
    c = Customer(business_id=b.id, first_name="F", last_name="X",
                 email="f@t.in", state_code=cust_state)
    db.add(c)
    db.commit()
    db.refresh(c)
    o = Order(business_id=b.id, internal_order_number=f"FF-{uniq}",
              shopify_order_id=f"gid://{uniq}", shopify_order_name=name,
              customer_id=c.id, currency="INR",
              subtotal_amount=float(total) - float(tax),
              discount_amount=0, shipping_amount=0, tax_amount=float(tax),
              total_amount=float(total), payment_status="PAID",
              financial_status="PAID", fulfillment_status="UNFULFILLED",
              operational_status="NEW", order_date=dt,
              ship_state_code=buyer_state, place_of_supply=buyer_state,
              business_state_code=seller_state)
    db.add(o)
    db.commit()
    db.refresh(o)
    ls.record_event(db, b.id, "SALE", total, tax, transaction_date=dt,
                    order_id=o.id, reference_number=name,
                    idempotency_key=f"FF-SALE:{o.id}", tally_voucher_type="Sales",
                    tally_voucher_number=name)
    db.commit()
    return o


# --- F1: close_checks fail-closed -------------------------------------------

def test_f1_check_query_failure_blocks_close():
    from app.services import accounting_service as acct
    mk, db, b = _mkdb()

    class _Boom:
        def query(self, *a, **k):
            raise RuntimeError("simulated DB outage")

    checks = acct.close_checks(_Boom(), b.id, 2026, 9)
    assert checks["total"] > 0
    assert len(checks["checks"]["check_errors"]) > 0
    try:
        acct.close_month(_Boom(), b.id, 2026, 9, "u1")
        assert False, "dirty/unverifiable month must not close"
    except acct.CloseBlocked as e:
        assert e.issues["total"] > 0
    db.close()


def test_f1_close_api_returns_422_on_db_failure():
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import Session as _Session
    from app.database import Base, get_db
    import app.models  # noqa: F401
    from app.main import app
    from app.models.business import Business
    from app.models.user import User
    from app.services.auth_service import hash_password
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    Base.metadata.create_all(eng)
    mk = sessionmaker(bind=eng)
    db = mk()
    b = Business(name="F1B", email="f1b@t.in")
    db.add(b)
    db.commit()
    db.refresh(b)
    u = User(business_id=b.id, name="A", email="f1a@t.in",
             password_hash=hash_password("x"), role="ADMIN")
    db.add(u)
    db.commit()
    db.close()
    app.dependency_overrides[get_db] = lambda: mk()
    c = TestClient(app)
    tok = c.post("/api/v1/auth/login",
                 json={"email": "f1a@t.in", "password": "x"}).json()["data"]["token"]
    h = {"Authorization": f"Bearer {tok}"}
    orig_query = _Session.query
    try:
        def _boom(self, *a, **k):
            raise RuntimeError("simulated check-query failure")
        _Session.query = _boom
        r = c.post("/api/v1/accounting/periods/2026/9/close", headers=h)
        assert r.status_code == 422, r.text
        body = r.json()
        assert body["success"] is False
        assert body["error"]["code"] == "CLOSE_BLOCKED"
        assert body["error"]["issues"]["total"] > 0
    finally:
        _Session.query = orig_query
        app.dependency_overrides.clear()


# --- F2: jurisdiction-aware GST split ----------------------------------------

def test_f2_intra_state_cgst_sgst():
    from app.services import tally_service as ts
    mk, db, b = _mkdb()
    o = _mksale(db, b, "#INTRA", seller_state="MH", buyer_state="MH")
    rows = ts.collect_export_rows(db, b.id)
    sale = [r for r in rows if r["order_ref"] == "#INTRA" and r["voucher_type"] == "Sales"]
    assert sale and sale[0]["jurisdiction"] == "SAME_STATE", sale
    cg, sg, ig = ts._split_gst(180.0, sale[0]["jurisdiction"])
    assert (cg, sg, ig) == (90.0, 90.0, 0.0)
    v = ts.validate_export(db, b.id)
    assert not any(w["code"] == "IGST_UNVERIFIED" for w in v["warnings"]), v["warnings"]
    db.close()


def test_f2_inter_state_full_igst():
    from app.services import tally_service as ts
    mk, db, b = _mkdb()
    _mksale(db, b, "#INTER", seller_state="MH", buyer_state="DL")
    rows = ts.collect_export_rows(db, b.id)
    sale = [r for r in rows if r["order_ref"] == "#INTER" and r["voucher_type"] == "Sales"]
    assert sale and sale[0]["jurisdiction"] == "INTER_STATE", sale
    cg, sg, ig = ts._split_gst(180.0, sale[0]["jurisdiction"])
    assert (cg, sg, ig) == (0.0, 0.0, 180.0)
    out = ts.generate_workbook_export(db, b.id, "u1")
    from openpyxl import load_workbook
    import io
    ws = load_workbook(filename=io.BytesIO(out["content"]))["Sales"]
    data = [r for r in ws.iter_rows(min_row=2) if r[1].value == "#INTER"]
    assert data, "inter-state sale must be in workbook"
    assert data[0][8].value == 0 and data[0][9].value == 0 and data[0][10].value == 180.0
    db.close()


def test_f2_unknown_jurisdiction_warns_igst_unverified():
    from app.services import tally_service as ts
    mk, db, b = _mkdb()
    _mksale(db, b, "#UNK")  # no states anywhere
    rows = ts.collect_export_rows(db, b.id)
    sale = [r for r in rows if r["order_ref"] == "#UNK" and r["voucher_type"] == "Sales"]
    assert sale and sale[0]["jurisdiction"] == "UNKNOWN", sale
    v = ts.validate_export(db, b.id)
    codes = {w["code"] for w in v["warnings"]}
    assert "IGST_UNVERIFIED" in codes, codes
    assert "CA_REVIEW" in codes, codes
    # default split still CGST/SGST but flagged, never silently filed
    cg, sg, ig = ts._split_gst(180.0, "UNKNOWN")
    assert (cg, sg, ig) == (90.0, 90.0, 0.0)
    db.close()


# --- F3: IST date discipline --------------------------------------------------

def test_f3_evening_utc_lands_on_ist_date():
    from app.services import tally_service as ts
    # 2026-09-10 19:00 UTC == 2026-09-11 00:30 IST
    d = ts._xl_date(datetime(2026, 9, 10, 19, 0, tzinfo=timezone.utc))
    assert (d.year, d.month, d.day) == (2026, 9, 11), d
    assert d.tzinfo is None
    from app.services import report_service as rs
    d2 = rs._xldate("2026-09-10T19:00:00+00:00")
    assert (d2.year, d2.month, d2.day) == (2026, 9, 11), d2


# --- F4: strict timestamp -----------------------------------------------------

def test_f4_missing_or_invalid_timestamp_fails():
    from app.services import shipsagar_service as ss
    assert ss.verify_timestamp(None) is False
    assert ss.verify_timestamp("") is False
    assert ss.verify_timestamp("not-a-time") is False
    assert ss.verify_timestamp("1") is False  # ancient epoch -> stale
    from datetime import timezone as _tz
    fresh = str(datetime.now(_tz.utc).timestamp())
    assert ss.verify_timestamp(fresh) is True


def test_f4_webhook_rejects_missing_timestamp(monkeypatch):
    import hashlib
    import hmac as _hmac
    import json as _json
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from app.database import Base, get_db
    import app.models  # noqa: F401
    from app.main import app
    from app.models.business import Business
    from app.models.shipment import Shipment
    from app import config
    SECRET = "f4-secret"
    monkeypatch.setattr(config.settings, "shipsagar_webhook_secret", SECRET)
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    Base.metadata.create_all(eng)
    mk = sessionmaker(bind=eng)
    db = mk()
    b = Business(name="F4", email="f4@t.in")
    db.add(b)
    db.commit()
    db.refresh(b)
    s = Shipment(business_id=b.id, order_id="o1", parcel_id="p1",
                 carrier_code="INDIA_POST", awb_number="EM123456789IN",
                 tracking_status="READY_TO_SHIP")
    db.add(s)
    db.commit()
    bid = b.id
    db.close()
    app.dependency_overrides[get_db] = lambda: mk()
    try:
        c = TestClient(app)
        body = {"event_id": "f4-e1", "tracking_number": "EM123456789IN",
                "courier": "INDIA_POST", "status": "delivered"}
        raw = _json.dumps(body).encode()
        sig = _hmac.new(SECRET.encode(), raw, hashlib.sha256).hexdigest()
        r = c.post("/api/webhooks/shipsagar", content=raw, headers={
            "Content-Type": "application/json",
            "X-ShipSagar-Signature": sig,
            "X-Business-Id": bid})  # no timestamp header
        assert r.status_code == 401, r.text
        assert r.json()["error"]["code"] == "STALE_TIMESTAMP"
    finally:
        app.dependency_overrides.clear()


# --- F5: Decimal money math ----------------------------------------------------

def test_f5_no_float_drift_in_monthly_money():
    from app.models.order import Order
    mk, db, b = _mkdb()
    for i, amt in enumerate(("0.10", "0.20", "0.30")):
        o = Order(business_id=b.id, internal_order_number=f"F5-{i}",
                  shopify_order_id=f"F5-{i}", shopify_order_name=f"#F5-{i}",
                  order_date=datetime(2026, 9, 5, tzinfo=timezone.utc),
                  total_amount=amt, tax_amount="0.00")
        db.add(o)
    db.commit()
    from app.services import report_service as rs
    rep = rs.monthly_report(db, b.id, "2026-09")
    assert rep["money"]["gross"] == 0.6, rep["money"]
    assert rep["profitability"]["gross"] == 0.6
    db.close()


def test_f5_gst_split_quantized():
    from app.services import tally_service as ts
    cg, sg, ig = ts._split_gst(100.0, "SAME_STATE")
    assert cg + sg + ig == 100.0
    cg, sg, ig = ts._split_gst(0.07, "SAME_STATE")
    assert round(cg + sg + ig, 2) == 0.07
