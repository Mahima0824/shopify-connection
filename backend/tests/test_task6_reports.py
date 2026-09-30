"""Task 6: dashboards/GST/reports hardening + backfill. Plan #84/#85 exact numbers."""
from datetime import datetime, timezone

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
import app.models  # noqa: F401 - register all models
from app.main import app
from app.database import get_db


def _env():
    from app.models.business import Business
    from app.models.user import User
    from app.models.order import Order
    from app.services.auth_service import hash_password
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(eng)
    mk = sessionmaker(bind=eng)
    db = mk()
    b = Business(name="B", email="b@t.in", state_code="MH")
    db.add(b)
    db.commit()
    db.refresh(b)
    u = User(business_id=b.id, name="A", email="a@t.in", password_hash=hash_password("x"), role="ADMIN")
    db.add(u)
    db.commit()
    db.refresh(u)
    o = Order(business_id=b.id, internal_order_number="ORD-84", shopify_order_id="gid://84",
              shopify_order_name="#84", currency="INR", subtotal_amount=1000.0,
              discount_amount=0.0, shipping_amount=0.0, tax_amount=180.0,
              total_amount=1180.0, financial_status="PAID", operational_status="DELIVERED",
              ship_state_code="MH", place_of_supply="MH", business_state_code="MH",
              order_date=datetime(2026, 9, 10, tzinfo=timezone.utc))
    db.add(o)
    db.commit()
    db.refresh(o)
    from app.models.order import OrderItem
    db.add(OrderItem(business_id=b.id, order_id=o.id, title="Widget", sku="W-1",
                     quantity=1, price=1180.0, hsn_code="1234", gst_rate=18.0,
                     taxable_amount=1000.0, cgst_amount=90.0, sgst_amount=90.0,
                     igst_amount=0.0, tax_amount=180.0))
    db.commit()
    db.close()
    app.dependency_overrides[get_db] = lambda: mk()
    c = TestClient(app)
    tok = c.post("/api/v1/auth/login", json={"email": "a@t.in", "password": "x"}).json()["data"]["token"]
    return c, {"Authorization": f"Bearer {tok}"}, mk


RANGE = {"from": "2026-09-01T00:00:00+00:00", "to": "2026-10-01T00:00:00+00:00"}


def _mk84_ledger(c, h, order_id=None):
    events = [("SALE", "1180.00", "180.00"), ("COGS", "500.00", "0.00"),
              ("SHIPPING_EXPENSE", "70.00", "0.00"),
              ("PACKAGING_EXPENSE", "15.00", "0.00"),
              ("PAYMENT_GATEWAY_FEE", "20.00", "0.00")]
    for typ, amt, tax in events:
        body = {"transaction_type": typ, "amount": amt, "tax_amount": tax,
                "transaction_date": "2026-09-10T10:00:00+00:00",
                "debit_account": "D", "credit_account": "C"}
        if order_id:
            body["order_id"] = order_id  # live hooks always link events to the order
        r = c.post("/api/v1/ledger", json=body, headers=h)
        assert r.status_code == 200, r.text


def test_plan84_exact_numbers_via_profit_report():
    c, h, mk = _env()
    db = mk()
    from app.models.order import Order
    oid = db.query(Order).filter_by(shopify_order_name="#84").first().id
    db.close()
    _mk84_ledger(c, h, order_id=oid)
    p = c.get("/api/v1/reports/profit", params=RANGE,
              headers=h).json()["data"]["profit"]
    assert p["gross_profit"] == "500.00"      # 1000 - 500
    assert p["operating_profit"] == "395.00"  # 500 - 70 - 15 - 20
    assert p["label"] == "OPERATING PROFIT"
    assert p["margin_pct"] == "39.50"         # 395/1000*100
    assert p["warning"] == ""


def test_plan85_half_refund():
    c, h, mk = _env()
    db = mk()
    from app.models.order import Order
    oid = db.query(Order).filter_by(shopify_order_name="#84").first().id
    db.close()
    c.post("/api/v1/ledger",
           json={"transaction_type": "SALE", "amount": "1180.00", "tax_amount": "180.00",
                 "transaction_date": "2026-09-10T10:00:00+00:00", "order_id": oid,
                 "debit_account": "D", "credit_account": "C"}, headers=h)
    c.post("/api/v1/ledger",
           json={"transaction_type": "REFUND", "amount": "590.00", "tax_amount": "90.00",
                 "transaction_date": "2026-09-12T10:00:00+00:00", "order_id": oid,
                 "debit_account": "D", "credit_account": "C"}, headers=h)
    rev = c.get("/api/v1/reports/sales", params=RANGE,
                headers=h).json()["data"]["revenue"]
    assert rev["net_inclusive"] == "590.00"
    assert rev["net_exclusive"] == "500.00"   # taxable refund 500, GST refund 90


def test_fy_boundaries_apr1():
    from app.services.report_service import resolve_preset
    s, e = resolve_preset("financial_year", now=datetime(2026, 9, 30, tzinfo=timezone.utc))
    assert (s.year, s.month, s.day) == (2026, 4, 1)
    assert (e.year, e.month, e.day) == (2027, 4, 1)
    s2, e2 = resolve_preset("financial_year", now=datetime(2026, 3, 31, tzinfo=timezone.utc))
    assert (s2.year, s2.month) == (2025, 4)
    ls, le = resolve_preset("last_fy", now=datetime(2026, 9, 30, tzinfo=timezone.utc))
    assert (ls.year, ls.month, ls.day) == (2025, 4, 1)
    assert (le.year, le.month, le.day) == (2026, 4, 1)


def test_gst_same_state_ok_interstate_flagged():
    c, h, mk = _env()
    db = mk()
    from app.models.order import Order
    o = db.query(Order).filter_by(shopify_order_name="#84").first()
    oid = o.id
    db.close()
    ok = c.get("/api/v1/reports/gst/validate", params={"order_id": oid}, headers=h).json()["data"]
    assert ok["valid"] is True
    assert ok["jurisdiction"] == "SAME_STATE"
    assert any(w["code"] == "CA_REVIEW" for w in ok["warnings"])
    # Flip to inter-state buyer: CGST/SGST must now be flagged.
    db = mk()
    o = db.query(Order).filter_by(id=oid).first()
    o.place_of_supply = "DL"
    o.ship_state_code = "DL"
    db.commit()
    db.close()
    bad = c.get("/api/v1/reports/gst/validate", params={"order_id": oid}, headers=h).json()["data"]
    assert bad["jurisdiction"] == "INTER_STATE"
    assert bad["valid"] is False
    assert any(e["code"] == "GST_JURISDICTION_MISMATCH" for e in bad["errors"])


def test_duplicate_delivery_creates_exactly_one_sale():
    """#94: duplicate Shopify delivery must not duplicate the order or its SALE."""
    from app.services.shopify_service import upsert_order
    from app.services import ledger_service as ls
    from app.models.order import Order
    from app.models.financial_transaction import FinancialTransaction
    c, h, mk = _env()
    db = mk()
    from app.models.business import Business
    b = db.query(Business).first()
    payload = {"id": 999001, "name": "#DUP1", "total_price": "1180.00",
               "financial_status": "paid", "fulfillment_status": "unfulfilled",
               "created_at": "2026-09-10T10:00:00Z"}
    id1 = upsert_order(db, b.id, payload)
    db.commit()
    o1 = db.query(Order).filter_by(business_id=b.id, shopify_order_id="999001").first()
    assert o1 is not None
    ls.record_sale_from_order(db, b.id, o1)
    db.commit()
    id2 = upsert_order(db, b.id, payload)  # duplicate delivery
    db.commit()
    o2 = db.query(Order).filter_by(id=o1.id).first()
    ls.record_sale_from_order(db, b.id, o2)
    db.commit()
    assert id1 == id2
    assert db.query(Order).filter_by(business_id=b.id, shopify_order_id="999001").count() == 1
    assert db.query(FinancialTransaction).filter_by(
        business_id=b.id, order_id=o1.id, transaction_type="SALE").count() == 1
    db.close()
    r1 = c.post("/api/v1/ledger/backfill", headers=h).json()["data"]
    assert r1["total_created"] >= 1  # pre-existing #84 order still gets its events
    r2 = c.post("/api/v1/ledger/backfill", headers=h).json()["data"]
    assert r2["total_created"] == 0


def test_backfill_skips_cancelled_unpaid_sale():
    """#43/#100: cancelled before payment -> no SALE; cancelled after payment -> SALE stands."""
    from app.models.business import Business
    from app.models.order import Order
    from app.models.payment import Payment
    from app.models.financial_transaction import FinancialTransaction
    c, h, mk = _env()
    db = mk()
    b = db.query(Business).first()
    now = datetime(2026, 9, 11, tzinfo=timezone.utc)
    o_unpaid = Order(business_id=b.id, internal_order_number="ORD-CXU",
                     shopify_order_id="gid://cxu", shopify_order_name="#CXU",
                     currency="INR", total_amount=500.0, tax_amount=0.0,
                     financial_status="CANCELLED", cancelled_at=now, order_date=now)
    o_paid = Order(business_id=b.id, internal_order_number="ORD-CXP",
                   shopify_order_id="gid://cxp", shopify_order_name="#CXP",
                   currency="INR", total_amount=500.0, tax_amount=0.0,
                   financial_status="CANCELLED", cancelled_at=now, order_date=now)
    db.add_all([o_unpaid, o_paid])
    db.commit()
    db.refresh(o_unpaid)
    db.refresh(o_paid)
    db.add(Payment(business_id=b.id, order_id=o_paid.id, amount=500.0,
                   payment_status="PAID", method="ONLINE"))
    db.commit()
    unpaid_id, paid_id = o_unpaid.id, o_paid.id
    bid = b.id
    db.close()
    c.post("/api/v1/ledger/backfill", headers=h)
    db = mk()
    assert db.query(FinancialTransaction).filter_by(
        business_id=bid, order_id=unpaid_id, transaction_type="SALE").count() == 0
    assert db.query(FinancialTransaction).filter_by(
        business_id=bid, order_id=paid_id, transaction_type="SALE").count() == 1
    db.close()


def test_filtered_totals_match_filtered_rows():
    """Filtered revenue/profit must come from the same filtered order set, not the whole period."""
    from app.models.business import Business
    from app.models.order import Order
    c, h, mk = _env()
    db = mk()
    b = db.query(Business).first()
    o84 = db.query(Order).filter_by(shopify_order_name="#84").first()
    now = datetime(2026, 9, 12, tzinfo=timezone.utc)
    o_dl = Order(business_id=b.id, internal_order_number="ORD-DL1",
                 shopify_order_id="gid://dl1", shopify_order_name="#DL1",
                 currency="INR", subtotal_amount=500.0, tax_amount=90.0,
                 total_amount=590.0, financial_status="PAID",
                 ship_state_code="DL", place_of_supply="DL",
                 business_state_code="MH", order_date=now)
    db.add(o_dl)
    db.commit()
    db.refresh(o_dl)
    o84_id, dl_id = o84.id, o_dl.id
    db.close()
    for oid, amt, tax in ((o84_id, "1180.00", "180.00"), (dl_id, "590.00", "90.00")):
        r = c.post("/api/v1/ledger",
                   json={"transaction_type": "SALE", "amount": amt, "tax_amount": tax,
                         "transaction_date": "2026-09-12T10:00:00+00:00",
                         "order_id": oid, "debit_account": "D", "credit_account": "C"},
                   headers=h)
        assert r.status_code == 200, r.text
    all_rev = c.get("/api/v1/reports/sales", params=RANGE, headers=h).json()["data"]
    assert all_rev["orders"] == 2
    assert all_rev["revenue"]["gross_inclusive"] == "1770.00"
    dl = c.get("/api/v1/reports/sales", params={**RANGE, "state": "DL"}, headers=h).json()["data"]
    assert dl["orders"] == 1
    assert dl["revenue"]["gross_inclusive"] == "590.00"
    assert dl["revenue"]["net_exclusive"] == "500.00"
    mh = c.get("/api/v1/reports/sales", params={**RANGE, "state": "MH"}, headers=h).json()["data"]
    assert mh["orders"] == 1
    assert mh["revenue"]["gross_inclusive"] == "1180.00"


def test_xlsx_cells_are_numeric_and_dates_with_totals():
    """#79: money as numbers, dates as Excel dates, totals row with formulas."""
    import io
    from openpyxl import load_workbook
    c, h, _mk = _env()
    r = c.get("/api/v1/reports/sales", params={**RANGE, "format": "xlsx"}, headers=h)
    assert r.status_code == 200
    wb = load_workbook(filename=io.BytesIO(r.content))
    ws = wb["Sales"]
    total_cell = ws.cell(row=2, column=5).value
    assert isinstance(total_cell, (int, float)), total_cell
    assert abs(total_cell - 1180.0) < 0.01
    assert isinstance(ws.cell(row=2, column=2).value, datetime)
    last = ws.max_row
    assert ws.cell(row=last, column=1).value == "TOTAL"
    formula = ws.cell(row=last, column=5).value
    assert isinstance(formula, str) and formula.startswith("=SUM(")


def test_dashboard_kpis_and_incomplete_cost_warning():
    c, h, _mk = _env()
    d = c.get("/api/v1/reports/dashboard", params=RANGE,
              headers=h).json()["data"]["kpis"]
    for k in ("orders_total", "gross_sales", "taxable_sales", "gst", "net_sales",
              "payments_received", "payments_pending", "cogs", "gross_profit",
              "operating_profit", "profit_margin_pct", "delivered_shipments"):
        assert k in d, k
    # No COGS event yet -> warning must be present, never silent.
    assert d["profit_warning"] != ""
    assert "ESTIMATED" in d["profit_label"]


def test_report_filters_and_csv_export():
    c, h, _mk = _env()
    assert c.get("/api/v1/reports/sales", params={**RANGE, "status": "PAID"}, headers=h).json()["data"]["orders"] == 1
    assert c.get("/api/v1/reports/sales", params={**RANGE, "status": "CANCELLED"}, headers=h).json()["data"]["orders"] == 0
    assert c.get("/api/v1/reports/sales", params={**RANGE, "state": "MH"}, headers=h).json()["data"]["orders"] == 1
    assert c.get("/api/v1/reports/sales", params={**RANGE, "state": "DL"}, headers=h).json()["data"]["orders"] == 0
    r = c.get("/api/v1/reports/sales", params={**RANGE, "format": "csv"}, headers=h)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert "Order" in r.text
    x = c.get("/api/v1/reports/sales", params={**RANGE, "format": "xlsx"}, headers=h)
    assert x.status_code == 200
    assert "spreadsheetml" in x.headers["content-type"]
