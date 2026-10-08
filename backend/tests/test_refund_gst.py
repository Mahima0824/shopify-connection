"""Task 3: cancellation is a marker; refunds reverse revenue AND post GST legs."""
from datetime import datetime, timezone
from decimal import Decimal as _D

DT = datetime(2026, 9, 10, 10, 0, tzinfo=timezone.utc)


def _env(items=(), total="1180.00", tax="180.00"):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool
    from app.database import Base
    import app.models  # noqa: F401
    from app.models.business import Business
    from app.models.customer import Customer
    from app.models.order import Order, OrderItem
    eng = create_engine("sqlite://", connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    Base.metadata.create_all(eng)
    db = sessionmaker(bind=eng)()
    b = Business(name="B", email="b@t.in")
    db.add(b); db.commit(); db.refresh(b)
    c = Customer(business_id=b.id, first_name="C", last_name="T", email="c@t.in")
    db.add(c); db.commit(); db.refresh(c)
    o = Order(business_id=b.id, internal_order_number="ORD-1", shopify_order_id="g1",
              shopify_order_name="#1001", customer_id=c.id, currency="INR",
              subtotal_amount=float(total) - float(tax), discount_amount=0,
              shipping_amount=0, tax_amount=float(tax), total_amount=float(total),
              payment_status="PAID", financial_status="PAID",
              fulfillment_status="FULFILLED", operational_status="SHIPPED", order_date=DT)
    db.add(o); db.commit(); db.refresh(o)
    for idx, it in enumerate(items or ((None, "1000.00", "90.00", "90.00", "0.00", 1, 18.0),)):
        hsn, taxable, cgst, sgst, igst, qty, rate = it
        db.add(OrderItem(business_id=b.id, order_id=o.id, title=f"P{idx}", sku=f"S{idx}",
                         quantity=qty, price="1000.00", hsn_code=hsn,
                         taxable_amount=float(taxable), cgst_amount=float(cgst),
                         sgst_amount=float(sgst), igst_amount=float(igst),
                         gst_rate=rate, tax_amount=float(cgst) + float(sgst) + float(igst)))
    db.commit()
    return db, b, o


def _refund(db, b, o, amount):
    from app.models.refund import Refund
    r = Refund(business_id=b.id, order_id=o.id, shopify_refund_id="rf1",
               amount=float(amount), currency="INR", status="COMPLETED")
    db.add(r); db.commit(); db.refresh(r)
    return r


def test_refund_posts_real_cgst_sgst_legs():
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    db, b, o = _env()
    r = _refund(db, b, o, "1180.00")
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    legs = db.query(FinancialTransaction).filter_by(transaction_type="REFUND").all()
    assert len(legs) == 3, [(l.debit_account, l.credit_account, str(l.amount)) for l in legs]
    assert str(sum((l.amount for l in legs if l.pnl_class == "REVERSAL"), 0)) == "1180.00"


def test_refund_mixed_rates_is_not_a_flat_half_split():
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    items = [(None, "400.00", "10.00", "10.00", "0.00", 1, 5.0),
             (None, "600.00", "36.00", "36.00", "0.00", 1, 12.0),
             (None, "0.00", "27.00", "27.00", "34.00", 1, 18.0)]
    db, b, o = _env(items=items, total="1000.00", tax="180.00")
    r = _refund(db, b, o, "1000.00")
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    legs = db.query(FinancialTransaction).filter_by(transaction_type="REFUND").all()
    by = {(l.debit_account, l.credit_account): str(l.amount) for l in legs}
    assert by[("Sales Revenue", "Customer / Receivable")] == "820.00"
    assert by[("Output CGST", "Customer / Receivable")] == "73.00"
    assert by[("Output SGST", "Customer / Receivable")] == "73.00"
    assert by[("Output IGST", "Customer / Receivable")] == "34.00"


def test_refund_posts_a_separate_payout_voucher():
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    db, b, o = _env()
    r = _refund(db, b, o, "590.00")
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    pay = db.query(FinancialTransaction).filter_by(transaction_type="REFUND_PAYOUT").all()
    assert len(pay) == 1
    assert pay[0].debit_account == "Customer / Receivable"
    assert pay[0].credit_account == "Payment Gateway"
    assert str(pay[0].amount) == "590.00"
    assert pay[0].pnl_class == "NONE"


def test_paid_cancel_then_refund_nets_to_zero_not_negative():
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    db, b, o = _env()
    ls.record_event(db, b.id, "SALE", "1180.00", "180.00", transaction_date=DT,
                    order_id=o.id, idempotency_key=f"SALE:{o.id}")
    r = _refund(db, b, o, "1180.00")
    ls.record_cancellation(db, b.id, o)
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    from app.services.ledger_service import revenue_summary
    rows = [{"transaction_type": x.transaction_type, "amount": str(x.amount),
             "tax_amount": str(x.tax_amount), "pnl_class": x.pnl_class}
            for x in db.query(FinancialTransaction).all()]
    rev = revenue_summary(rows)
    assert rev["gross_inclusive"] == "1180.00"
    assert rev["refunds_inclusive"] == "1180.00"
    assert rev["net_inclusive"] == "0.00"
    # GST must net to zero too. The credit note's CGST/SGST legs carry the tax in
    # tax_amount, so revenue_summary subtracts it; posting it as amount alone left
    # gst at a full 180.00 and net_exclusive at -180.00 on a fully refunded order.
    assert rev["gst"] == "0.00"
    assert rev["net_exclusive"] == "0.00"


def test_cancellation_marker_carries_no_amount():
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    db, b, o = _env()
    ls.record_cancellation(db, b.id, o)
    db.commit()
    rows = db.query(FinancialTransaction).filter_by(transaction_type="CANCELLATION").all()
    assert len(rows) == 1
    assert str(rows[0].amount) == "0.00"
    assert str(rows[0].tax_amount) == "0.00"
    assert rows[0].pnl_class == "NONE"


def test_refund_credit_note_legs_sum_to_the_refund_exactly():
    """The taxable leg is derived from the tax legs, and post_voucher rounds each
    leg on its own. An unrounded split therefore left the credit note a paisa
    short of the refund on any amount that is not a clean fraction of the order.
    """
    from decimal import Decimal
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    for amount in ("1180.00", "590.00", "100.00", "333.33", "1.00", "777.77"):
        db, b, o = _env()
        r = _refund(db, b, o, amount)
        ls.record_refund_from_refund(db, b.id, r, order=o)
        db.commit()
        legs = db.query(FinancialTransaction).filter_by(
            refund_id=r.id, transaction_type="REFUND").all()
        total = sum((l.amount for l in legs), Decimal("0"))
        assert total == r.amount, f"{amount}: legs {[str(l.amount) for l in legs]} sum {total}"


def test_refund_legs_do_not_collide_on_one_voucher_number():
    """One refund writes several vouchers, so they must not share a
    reference_number: tally_service derives voucher_number from it, and a repeat
    blocks the whole export with VOUCHER_NOT_UNIQUE."""
    from app.services import tally_service as ts
    from app.services import ledger_service as ls
    db, b, o = _env()
    ls.record_sale_from_order(db, b.id, o)
    r = _refund(db, b, o, "1180.00")
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    nums = [row["voucher_number"] for row in ts.collect_export_rows(db, b.id)]
    assert len(nums) == len(set(nums)), nums
    assert ts.validate_export(db, b.id)["can_export"] is True


def test_split_sale_voucher_passes_export_validation():
    """The validator compares a voucher's LEGS against the order total, so any
    multi-leg SALE failed INVOICE_TOTAL_MISMATCH on every leg."""
    from app.services import ledger_service as ls
    from app.services import tally_service as ts
    db, b, o = _env()
    ls.record_sale_from_order(db, b.id, o)
    db.commit()
    rows = ts.collect_export_rows(db, b.id)
    assert len(rows) == 3, [(r["debit_account"], r["amount"]) for r in rows]
    assert all(r["voucher_group_id"] for r in rows), "voucher_group_id must be projected"
    v = ts.validate_export(db, b.id)
    assert v["can_export"] is True, v["errors"]


def test_split_sale_reports_a_wrong_total_once_not_per_leg():
    """A voucher's legs are compared together, so a bad invoice is one error."""
    from app.services import ledger_service as ls
    from app.services import tally_service as ts
    db, b, o = _env()
    ls.record_sale_from_order(db, b.id, o)
    db.commit()
    from app.models.financial_transaction import FinancialTransaction
    legs = db.query(FinancialTransaction).filter_by(transaction_type="SALE").all()
    legs[-1].amount = _D("1.00")  # voucher's legs no longer total the order
    db.commit()
    v = ts.validate_export(db, b.id)
    mism = [e for e in v["errors"] if e["code"] == "INVOICE_TOTAL_MISMATCH"]
    assert len(mism) == 1, mism
    assert "ledger" in mism[0]["message"] and "!=" in mism[0]["message"], mism[0]


def test_legacy_single_leg_sale_still_uses_the_per_row_check():
    """A NULL voucher_group_id row has no group to sum, so it keeps the old
    per-row comparison and must not be measured against an unrelated group."""
    from datetime import datetime, timezone
    from app.models.financial_transaction import FinancialTransaction
    from app.models.order import Order
    from app.services import ledger_service as ls
    from app.services import tally_service as ts
    db, b, o = _env()
    o.total_amount = _D("1180.00")
    db.add(FinancialTransaction(
        business_id=b.id, transaction_id="TXN-LEGACY-SALE", order_id=o.id,
        transaction_type="SALE", transaction_date=datetime(2026, 9, 10, 10, 0, tzinfo=timezone.utc),
        amount="1180.00", tax_amount="180.00", net_amount="1000.00", currency="INR",
        debit_account="Customer / Receivable", credit_account="Sales Revenue",
        idempotency_key=f"SALE:{o.id}", pnl_class="REVENUE", leg_index=0,
        voucher_group_id=None))
    # A second, unrelated grouped SALE whose voucher totals 2360.00. If the legacy
    # NULL-group row were ever measured against a group, or all rows summed, this
    # would break the assertion below -- with a single row both paths agree.
    other = Order(business_id=b.id, internal_order_number="ORD-2", shopify_order_id="g2",
                  shopify_order_name="#2002", customer_id=o.customer_id, currency="INR",
                  subtotal_amount=2000.0, discount_amount=0, shipping_amount=0,
                  tax_amount=360.0, total_amount=2360.0, payment_status="PAID",
                  financial_status="PAID", order_date=DT)
    db.add(other); db.commit(); db.refresh(other)
    ls.record_sale_from_order(db, b.id, other)
    db.commit()
    rows = ts.collect_export_rows(db, b.id)
    legacy = [r for r in rows if r["transaction_id"] == "TXN-LEGACY-SALE"]
    assert len(legacy) == 1 and legacy[0]["voucher_group_id"] is None
    assert ts.validate_export(db, b.id)["can_export"] is True


def test_legacy_sale_row_at_the_split_key_raises_loudly():
    """A pre-0020 SALE row holds SALE:{id} with no group. It cannot satisfy a
    3-leg voucher, so post_voucher must raise rather than append a second sale."""
    from datetime import datetime, timezone
    from app.models.financial_transaction import FinancialTransaction
    from app.services import ledger_service as ls
    from app.services import tally_service as ts
    db, b, o = _env()
    o.total_amount = _D("1180.00")
    db.add(FinancialTransaction(
        business_id=b.id, transaction_id="TXN-LEGACY-SALE", order_id=o.id,
        transaction_type="SALE", transaction_date=datetime(2026, 9, 10, 10, 0, tzinfo=timezone.utc),
        amount="1180.00", tax_amount="180.00", net_amount="1000.00", currency="INR",
        debit_account="Customer / Receivable", credit_account="Sales Revenue",
        idempotency_key=f"SALE:{o.id}", pnl_class="REVENUE", leg_index=0,
        voucher_group_id=None))
    db.commit()
    before = db.query(FinancialTransaction).count()
    raised = False
    try:
        ls.record_sale_from_order(db, b.id, o)
    except ls.VoucherKeyConflict:
        raised = True
    db.rollback()
    assert raised, "a legacy single-leg holder must not silently accept a multi-leg post"
    assert db.query(FinancialTransaction).count() == before, "no duplicate sale written"


def test_inter_state_sale_credits_igst_and_not_cgst_sgst():
    """An inter-state order's whole tax belongs on Output IGST."""
    from app.services import ledger_service as ls
    db, b, o = _env(items=((None, "1000.00", "0.00", "0.00", "180.00", 1, 18.0),),
                    total="1180.00", tax="180.00")
    ls.record_sale_from_order(db, b.id, o)
    db.commit()
    from app.models.financial_transaction import FinancialTransaction
    legs = db.query(FinancialTransaction).filter_by(transaction_type="SALE").all()
    credited = {l.credit_account: str(l.amount) for l in legs}
    assert credited == {"Sales Revenue": "1000.00", "Output IGST": "180.00"}, credited


def test_order_timeline_reports_the_stored_leg_not_a_re_derivation():
    """A stored leg IS the double entry. Re-deriving it through to_double_entry's
    legacy whole-sale branch gave each GST leg a zero-value credit plus a credit to
    'Output GST', a ledger that does not exist in the chart."""
    from app.services import ledger_service as ls
    db, b, o = _env()
    ls.record_sale_from_order(db, b.id, o)
    db.commit()
    items = ls.order_timeline(db, b.id, o.id)
    assert len(items) == 3, items
    chart = {ls.ACC_RECEIVABLE, ls.ACC_SALES, ls.ACC_OUTPUT_CGST,
             ls.ACC_OUTPUT_SGST, ls.ACC_OUTPUT_IGST}
    for it in items:
        assert it["double_entry"] == [{
            "debit": it["debit_account"],
            "credit": it["credit_account"],
            "amount": it["amount"],
        }], it
        for de in it["double_entry"]:
            assert de["debit"] in chart, de
            assert de["credit"] in chart, de
            assert _D(de["amount"]) > 0, f"zero-value leg rendered: {it}"
        assert it["pnl_class"] in {"REVENUE", "REVERSAL", "COGS", "EXPENSE", "NONE"}


def test_sales_items_join_picks_the_invoice_number_not_the_first_leg():
    """A sale is 3 rows sharing one transaction_date, so 'first row for this order'
    is arbitrary; only the taxable leg carries the invoice number and the GST legs
    fall back to a bare transaction_id. Fed the legs in the worst order, the item
    rows must still be tagged with the invoice number the Sales sheet shows."""
    import io
    from openpyxl import load_workbook
    from app.services import ledger_service as ls
    from app.services import tally_service as ts
    db, b, o = _env()
    ls.record_sale_from_order(db, b.id, o)
    db.commit()
    rows = ts.collect_export_rows(db, b.id)
    sales = [r for r in rows if r["voucher_type"] == "Sales"]
    assert len(sales) == 3, sales
    # Worst case: the bare-transaction_id GST legs are seen before the taxable leg,
    # so a plain first-wins setdefault would tag the item rows TXN-0000000N.
    worst = [r for r in sales if r["voucher_number"] == r["transaction_id"]]
    worst += [r for r in sales if r["voucher_number"] != r["transaction_id"]]
    content = ts.build_workbook(rows, ts.ledger_mapping_completeness(db, b.id)["mappings"],
                                {"valid": 1, "error_count": 0, "warning_count": 0,
                                 "can_export": True}, "Sep 2026", "BATCH")
    filled = ts._fill_sales_items(db, b.id, content, worst)
    ws = load_workbook(filename=io.BytesIO(filled))["Sales_Items"]
    items = [r for r in ws.iter_rows(min_row=2) if r[0].value not in (None, "TOTAL")]
    assert items, "expected item rows"
    assert {r[0].value for r in items} == {o.shopify_order_name}, [r[0].value for r in items]


def test_sales_items_invoice_matches_the_sales_sheet_end_to_end():
    import io
    from openpyxl import load_workbook
    from app.services import ledger_service as ls
    from app.services import tally_service as ts
    db, b, o = _env()
    ls.record_sale_from_order(db, b.id, o)
    db.commit()
    out = ts.generate_workbook_export(db, b.id, "u1")
    wb = load_workbook(filename=io.BytesIO(out["content"]))
    sales = [r for r in wb["Sales"].iter_rows(min_row=2) if r[0].value not in (None, "TOTAL")]
    items = [r for r in wb["Sales_Items"].iter_rows(min_row=2) if r[0].value not in (None, "TOTAL")]
    assert items, "expected item rows"
    assert o.shopify_order_name in {r[1].value for r in sales}
    assert {r[0].value for r in items} == {o.shopify_order_name}


def test_refund_replay_is_idempotent():
    """Webhook redelivery re-posts the same refund; it must not duplicate a leg
    or drop one."""
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    db, b, o = _env()
    r = _refund(db, b, o, "1180.00")
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    first = {(t.transaction_type, t.debit_account): t.amount
             for t in db.query(FinancialTransaction).all()}
    for _ in range(3):
        ls.record_refund_from_refund(db, b.id, r, order=o)
        db.commit()
    rows = db.query(FinancialTransaction).all()
    assert len(rows) == len(first), f"{len(first)} -> {len(rows)} rows"
    after = {(t.transaction_type, t.debit_account): t.amount for t in rows}
    assert after == first


def test_refund_falls_back_to_half_split_without_order_items():
    """No OrderItem rows carrying split tax: the tax still has to reach the credit
    note, so it is halved from the order total rather than dropped to zero."""
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    db, b, o = _env(items=((None, "1000.00", "0.00", "0.00", "0.00", 1, 18.0),))
    r = _refund(db, b, o, "1180.00")
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    legs = db.query(FinancialTransaction).filter_by(
        refund_id=r.id, transaction_type="REFUND").all()
    by = {l.debit_account: l.amount for l in legs}
    assert by["Output CGST"] == _D("90.00")
    assert by["Output SGST"] == _D("90.00")
    assert by["Sales Revenue"] == _D("1000.00")


def test_refund_posts_a_negative_taxable_leg_rather_than_dropping_it():
    """Order tax above its total makes the taxable leg negative. Dropping it left
    the receivable funded with no matching revenue reversal."""
    from app.services import ledger_service as ls
    from app.models.financial_transaction import FinancialTransaction
    db, b, o = _env(items=((None, "1000.00", "600.00", "600.00", "0.00", 1, 120.0),),
                    total="1000.00", tax="1200.00")
    r = _refund(db, b, o, "1000.00")
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    legs = db.query(FinancialTransaction).filter_by(
        refund_id=r.id, transaction_type="REFUND").all()
    taxable = [l for l in legs if l.debit_account == "Sales Revenue"]
    assert len(taxable) == 1, "the taxable leg must be posted, even when negative"
    assert taxable[0].amount == _D("-200.00")  # 1000.00 refund - 1200.00 tax


def test_backfill_refund_counts_do_not_inflate_on_rerun():
    """One refund now writes several vouchers keyed 'REFUND:{id}:taxable' etc., so
    an exact-match guard never fired: every run re-entered the writer, wrote
    nothing (post_voucher is idempotent) and still counted as created."""
    from app.services import ledger_service as ls
    from app.services import report_service as rs
    db, b, o = _env()
    r = _refund(db, b, o, "1180.00")
    ls.record_refund_from_refund(db, b.id, r, order=o)
    db.commit()
    first = rs.backfill_ledger(db, b.id)
    assert first["refunds"] == {"created": 0, "skipped": 1}, first
    second = rs.backfill_ledger(db, b.id)
    assert second["refunds"] == {"created": 0, "skipped": 1}, second


def test_backfill_still_skips_a_legacy_single_key_refund_row():
    """A pre-0020 row keyed exactly 'REFUND:{id}' must still match the guard."""
    from app.models.financial_transaction import FinancialTransaction
    from app.services import report_service as rs
    db, b, o = _env()
    r = _refund(db, b, o, "1180.00")
    legacy = FinancialTransaction(
        business_id=b.id, transaction_id="TXN-LEGACY", order_id=o.id,
        refund_id=r.id, transaction_type="REFUND", transaction_date=DT,
        amount="1180.00", tax_amount="180.00", net_amount="1000.00", currency="INR",
        debit_account="Sales Revenue", credit_account="Payment Gateway",
        idempotency_key=f"REFUND:{r.id}", pnl_class="REVERSAL", leg_index=0,
    )
    db.add(legacy)
    db.commit()
    before = db.query(FinancialTransaction).filter_by(refund_id=r.id).count()
    out = rs.backfill_ledger(db, b.id)
    after = db.query(FinancialTransaction).filter_by(refund_id=r.id).count()
    assert out["refunds"]["skipped"] == 1, out
    assert after == before, f"legacy row was rewritten: {before} -> {after}"
